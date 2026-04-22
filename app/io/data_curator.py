"""
Data Curation module — reads illuminated J-V xlsx files, applies thresholds,
and exports a curated workbook with Raw_Data, Curated_[param], Statistics,
and individual per-cell J-V sheets.

Column positions verified against app/io/exporter.py save_data():
  xlsxwriter 0-indexed col → openpyxl 1-indexed col (row 2 = scalar params)
  Voc_fwd H(8), Jsc_fwd I(9), FF_fwd J(10), PCE_fwd K(11), Rs_fwd L(12), Rsh_fwd M(13)
  Voc_rev N(14), Jsc_rev O(15), FF_rev P(16), PCE_rev Q(17), Rs_rev R(18), Rsh_rev S(19)
  HI T(20)
  Raw J-V data (rows 3+): col A=v_fwd, B=j_fwd, C=p_fwd, D=v_rev, E=j_rev, F=p_rev
"""

import io
import os
from pathlib import Path
from datetime import datetime

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

_SHEET = "MeasurementData"

_COL = {
    "Voc_fwd": 8,
    "Jsc_fwd": 9,
    "FF_fwd":  10,
    "PCE_fwd": 11,
    "Rs_fwd":  12,
    "Rsh_fwd": 13,
    "Voc_rev": 14,
    "Jsc_rev": 15,
    "FF_rev":  16,
    "PCE_rev": 17,
    "Rs_rev":  18,
    "Rsh_rev": 19,
    "HI":      20,
}

PARAMS      = ["Voc", "Jsc", "FF", "PCE", "HI", "Rs", "Rsh"]
P_HEADERS   = ["Voc (V)", "Jsc (mA/cm²)", "FF (%)", "PCE (%)",
               "HI (%)", "Rs (Ω·cm²)", "Rsh (Ω·cm²)"]
P_FMT       = ["0.0000", "0.0000", "0.00", "0.0000",
               "0.00",   "0.0000", "0.0"]


class CellData:
    """All PV parameters for one illuminated J-V measurement file."""

    def __init__(self, filepath):
        self.filepath  = filepath
        self.filename  = Path(filepath).name
        self.device    = self._extract_device()

        self.Voc_fwd = self.Jsc_fwd = self.FF_fwd = self.PCE_fwd = 0.0
        self.Rs_fwd  = self.Rsh_fwd = 0.0
        self.Voc_rev = self.Jsc_rev = self.FF_rev = self.PCE_rev = 0.0
        self.Rs_rev  = self.Rsh_rev = 0.0
        self.HI      = 0.0

    def _extract_device(self):
        stem  = Path(self.filepath).stem
        parts = stem.split("_")
        for i, p in enumerate(parts):
            if len(p) == 8 and p.isdigit():
                return "_".join(parts[:i]) or stem
        return parts[0] if parts else stem

    def get_val(self, param, scan):
        """Return value for param at the given scan direction ('fwd' or 'rev')."""
        if param == "HI":
            return self.HI
        return getattr(self, f"{param}_{scan}", 0.0)

    def passes(self, param, mn, mx, scan="rev"):
        """Return True if this cell passes [mn, mx] for param.

        scan='both' requires fwd AND rev each to be within [mn, mx].
        """
        if scan == "both":
            return (mn <= self.get_val(param, "fwd") <= mx and
                    mn <= self.get_val(param, "rev") <= mx)
        return mn <= self.get_val(param, scan) <= mx


def load_cell_data(filepath):
    """Read a MeasurementData xlsx file.  Returns CellData or None on failure."""
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        if _SHEET not in wb.sheetnames:
            wb.close()
            return None
        ws = wb[_SHEET]

        def _v(col):
            val = ws.cell(2, col).value
            return float(val) if val is not None else 0.0

        c = CellData(filepath)
        c.Voc_fwd = _v(_COL["Voc_fwd"])
        c.Jsc_fwd = _v(_COL["Jsc_fwd"])
        c.FF_fwd  = _v(_COL["FF_fwd"])
        c.PCE_fwd = _v(_COL["PCE_fwd"])
        c.Rs_fwd  = _v(_COL["Rs_fwd"])
        c.Rsh_fwd = _v(_COL["Rsh_fwd"])
        c.Voc_rev = _v(_COL["Voc_rev"])
        c.Jsc_rev = _v(_COL["Jsc_rev"])
        c.FF_rev  = _v(_COL["FF_rev"])
        c.PCE_rev = _v(_COL["PCE_rev"])
        c.Rs_rev  = _v(_COL["Rs_rev"])
        c.Rsh_rev = _v(_COL["Rsh_rev"])
        c.HI      = _v(_COL["HI"])
        wb.close()
        return c
    except Exception as exc:
        print(f"Warning: could not load {filepath}: {exc}")
        return None


def group_by_substrate(cells, n_per_substrate=4):
    """Group cells sequentially: cells 0..N-1 = Substrate_1, N..2N-1 = Substrate_2, etc.

    The last group may have fewer than n_per_substrate cells.
    """
    groups = {}
    for i, cell in enumerate(cells):
        substrate_num = i // n_per_substrate + 1
        key = f"Substrate_{substrate_num}"
        groups.setdefault(key, []).append(cell)
    return groups


def _create_jv_chart_image(v_fwd, j_fwd, v_rev, j_rev, device_name):
    """Generate a J-V chart and return it as a BytesIO PNG image."""
    fig, ax = plt.subplots(figsize=(6, 4))
    ax.plot(v_fwd, j_fwd, 'b-o', markersize=3, linewidth=1.5,
            label='Forward', markerfacecolor='blue')
    ax.plot(v_rev, j_rev, 'r--o', markersize=3, linewidth=1.5,
            label='Reverse', markerfacecolor='white', markeredgecolor='red')
    ax.set_xlabel('Voltage (V)', fontsize=11)
    ax.set_ylabel('Current Density (mA/cm²)', fontsize=11)
    ax.set_title(f'J-V Curve — {device_name}', fontsize=12, fontweight='bold')
    ax.legend(fontsize=10)
    ax.grid(True, color='#E0E0E0', linewidth=0.6)
    ax.set_facecolor('white')
    fig.patch.set_facecolor('white')
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=120, bbox_inches='tight')
    plt.close(fig)
    buf.seek(0)
    return buf


def _add_individual_sheet(wb, cell_data, used_names):
    """Add one sheet to wb with raw J-V data and an embedded J-V chart."""
    base_name  = cell_data.device[:28]
    sheet_name = base_name
    counter    = 2
    while sheet_name in used_names:
        sheet_name = f"{base_name}_{counter}"
        counter   += 1
    used_names.add(sheet_name)

    ws = wb.create_sheet(sheet_name)

    ws['A1'] = cell_data.device
    ws['A1'].font = Font(bold=True, size=13, color="1565C0")

    fwd_headers = ['Voltage (V)', 'J_fwd (mA/cm²)']
    rev_headers = ['Voltage (V)', 'J_rev (mA/cm²)']
    for k, h in enumerate(fwd_headers):
        c = ws.cell(3, k + 1, h)
        c.fill = PatternFill("solid", fgColor="1565C0")
        c.font = Font(bold=True, color="FFFFFF", size=10)
    for k, h in enumerate(rev_headers):
        c = ws.cell(3, k + 4, h)
        c.fill = PatternFill("solid", fgColor="B71C1C")
        c.font = Font(bold=True, color="FFFFFF", size=10)

    try:
        wb_src = openpyxl.load_workbook(cell_data.filepath, read_only=True, data_only=True)
        ws_src = wb_src[_SHEET]

        v_fwd, j_fwd, v_rev, j_rev = [], [], [], []
        data_row = 4

        for src_row in ws_src.iter_rows(min_row=3, values_only=True):
            # exporter.py layout: A=v_fwd(0), B=j_fwd(1), D=v_rev(3), E=j_rev(4)
            vf = src_row[0] if len(src_row) > 0 else None
            jf = src_row[1] if len(src_row) > 1 else None
            vr = src_row[3] if len(src_row) > 3 else None
            jr = src_row[4] if len(src_row) > 4 else None

            if vf is not None and isinstance(vf, (int, float)):
                c = ws.cell(data_row, 1, round(float(vf), 5))
                c.number_format = '0.00000'
                if jf is not None and isinstance(jf, (int, float)):
                    c2 = ws.cell(data_row, 2, round(float(jf), 5))
                    c2.number_format = '0.00000'
                    j_fwd.append(float(jf))
                v_fwd.append(float(vf))

            if vr is not None and isinstance(vr, (int, float)):
                c = ws.cell(data_row, 4, round(float(vr), 5))
                c.number_format = '0.00000'
                if jr is not None and isinstance(jr, (int, float)):
                    c2 = ws.cell(data_row, 5, round(float(jr), 5))
                    c2.number_format = '0.00000'
                    j_rev.append(float(jr))
                v_rev.append(float(vr))

            data_row += 1

        wb_src.close()

        if v_fwd and v_rev and j_fwd and j_rev:
            img_buf = _create_jv_chart_image(v_fwd, j_fwd, v_rev, j_rev, cell_data.device)
            img = XLImage(img_buf)
            img.anchor = 'G3'
            ws.add_image(img)

    except Exception as e:
        ws['A2'] = f"Could not read source data: {e}"
        ws['A2'].font = Font(color="C62828", italic=True)

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 3
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 18


def export_curated_excel(cells, thresholds, threshold_scan,
                         substrate_groups, output_path,
                         progress_callback=None):
    """Build and save the curated Excel workbook.

    Parameters
    ----------
    cells : list[CellData]
    thresholds : dict {param: (min_val, max_val)}  — only enabled params
    threshold_scan : 'rev' or 'both'
    substrate_groups : dict returned by group_by_substrate()
    output_path : full .xlsx path
    progress_callback : optional callable(str)
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ── Styles ───────────────────────────────────────────────────────────────
    H_FILL  = PatternFill("solid", fgColor="1565C0")
    H_FONT  = Font(bold=True, color="FFFFFF", size=10)
    SH_FILL = PatternFill("solid", fgColor="BBDEFB")
    SH_FONT = Font(bold=True, color="0D47A1", size=10)
    PASS_F  = PatternFill("solid", fgColor="C8E6C9")
    AVG_F   = PatternFill("solid", fgColor="FFF9C4")
    STD_F   = PatternFill("solid", fgColor="F3E5F5")
    B_FONT  = Font(bold=True, size=10)
    N_FONT  = Font(size=10)
    C_ALIGN = Alignment(horizontal="center", vertical="center")
    L_ALIGN = Alignment(horizontal="left",   vertical="center")

    PASS_HDRS = [f"Pass_{p}" for p in thresholds]

    # Column layout (1-indexed):
    # FWD: Device | 7 params | pass cols | GAP | REV: Device | 7 params | pass cols
    FWD_DEV  = 1
    FWD_P    = list(range(2, 9))
    FWD_PASS = list(range(9, 9 + len(thresholds)))
    GAP      = 9 + len(thresholds)
    REV_DEV  = GAP + 1
    REV_P    = list(range(REV_DEV + 1, REV_DEV + 8))
    REV_PASS = list(range(REV_DEV + 8, REV_DEV + 8 + len(thresholds)))

    def _write_headers(ws, start_row):
        end_fwd = (FWD_PASS[-1] if FWD_PASS else FWD_P[-1])
        end_rev = (REV_PASS[-1] if REV_PASS else REV_P[-1])

        ws.cell(start_row, FWD_DEV, "FORWARD SCAN").fill = SH_FILL
        ws.cell(start_row, FWD_DEV).font = SH_FONT
        ws.merge_cells(start_row=start_row, start_column=FWD_DEV,
                       end_row=start_row,   end_column=end_fwd)
        ws.cell(start_row, REV_DEV, "REVERSE SCAN").fill = SH_FILL
        ws.cell(start_row, REV_DEV).font = SH_FONT
        ws.merge_cells(start_row=start_row, start_column=REV_DEV,
                       end_row=start_row,   end_column=end_rev)
        r = start_row + 1
        for col, txt in [(FWD_DEV, "Device")] + list(zip(FWD_P, P_HEADERS)):
            ws.cell(r, col, txt).fill = H_FILL
            ws.cell(r, col).font = H_FONT
            ws.cell(r, col).alignment = C_ALIGN
        for k, ph in enumerate(PASS_HDRS):
            ws.cell(r, FWD_PASS[k], ph).fill = H_FILL
            ws.cell(r, FWD_PASS[k]).font = H_FONT
            ws.cell(r, FWD_PASS[k]).alignment = C_ALIGN
        for col, txt in [(REV_DEV, "Device")] + list(zip(REV_P, P_HEADERS)):
            ws.cell(r, col, txt).fill = H_FILL
            ws.cell(r, col).font = H_FONT
            ws.cell(r, col).alignment = C_ALIGN
        for k, ph in enumerate(PASS_HDRS):
            ws.cell(r, REV_PASS[k], ph).fill = H_FILL
            ws.cell(r, REV_PASS[k]).font = H_FONT
            ws.cell(r, REV_PASS[k]).alignment = C_ALIGN
        return r + 1

    def _write_cell_row(ws, row, cell, highlight=True):
        for side, dev_col, p_cols, pass_cols in [
            ("fwd", FWD_DEV, FWD_P, FWD_PASS),
            ("rev", REV_DEV, REV_P, REV_PASS),
        ]:
            ws.cell(row, dev_col, cell.device).font = N_FONT
            ws.cell(row, dev_col).alignment = L_ALIGN
            for k, (p, fmt) in enumerate(zip(PARAMS, P_FMT)):
                v = cell.get_val(p, side)
                ws.cell(row, p_cols[k], round(float(v), 6)).font = N_FONT
                ws.cell(row, p_cols[k]).number_format = fmt
            all_pass = True
            for k, (param, (mn, mx)) in enumerate(thresholds.items()):
                scan_for_check = side if threshold_scan == "both" else "rev"
                passed = cell.passes(param, mn, mx, scan=scan_for_check)
                sym = "✓" if passed else "✗"
                ws.cell(row, pass_cols[k], sym).font = Font(
                    color="2E7D32" if passed else "C62828", bold=True, size=10)
                ws.cell(row, pass_cols[k]).alignment = C_ALIGN
                if not passed:
                    all_pass = False
            if highlight and all_pass and thresholds:
                for col in [dev_col] + p_cols + list(pass_cols):
                    ws.cell(row, col).fill = PASS_F

    def _write_stats_both(ws, row, grp_cells):
        """Write Average and Std Dev rows with FWD and REV on the same row."""
        fwd_mat = [[c.get_val(p, "fwd") for p in PARAMS] for c in grp_cells]
        rev_mat = [[c.get_val(p, "rev") for p in PARAMS] for c in grp_cells]
        fwd_arr = np.array(fwd_mat, dtype=float)
        rev_arr = np.array(rev_mat, dtype=float)
        fwd_avgs = np.mean(fwd_arr, axis=0)
        fwd_stds = np.std(fwd_arr,  axis=0)
        rev_avgs = np.mean(rev_arr, axis=0)
        rev_stds = np.std(rev_arr,  axis=0)

        for label, fwd_vals, rev_vals, fill in [
            ("Average", fwd_avgs, rev_avgs, AVG_F),
            ("Std Dev", fwd_stds, rev_stds, STD_F),
        ]:
            ws.cell(row, FWD_DEV, label).font = B_FONT
            ws.cell(row, FWD_DEV).fill = fill
            for k, (v, fmt) in enumerate(zip(fwd_vals, P_FMT)):
                ws.cell(row, FWD_P[k], round(float(v), 4)).font = B_FONT
                ws.cell(row, FWD_P[k]).fill = fill
                ws.cell(row, FWD_P[k]).number_format = fmt

            ws.cell(row, REV_DEV, label).font = B_FONT
            ws.cell(row, REV_DEV).fill = fill
            for k, (v, fmt) in enumerate(zip(rev_vals, P_FMT)):
                ws.cell(row, REV_P[k], round(float(v), 4)).font = B_FONT
                ws.cell(row, REV_P[k]).fill = fill
                ws.cell(row, REV_P[k]).number_format = fmt

            row += 1
        return row

    def _autofit(ws):
        for col in ws.columns:
            mx = max((len(str(c.value or "")) for c in col), default=6)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(mx + 2, 10)
        ws.column_dimensions[get_column_letter(GAP)].width = 3

    # ── Sheet: Raw_Data ──────────────────────────────────────────────────────
    if progress_callback:
        progress_callback("Building Raw_Data sheet...")
    ws_raw = wb.create_sheet("Raw_Data")
    row = _write_headers(ws_raw, 1)
    for sub_name, sub_cells in substrate_groups.items():
        ws_raw.cell(row, FWD_DEV, f"── {sub_name} ──").font = Font(
            bold=True, italic=True, size=9, color="5D4037")
        row += 1
        for cell in sub_cells:
            _write_cell_row(ws_raw, row, cell)
            row += 1
        row = _write_stats_both(ws_raw, row, sub_cells)
        row += 1
    _autofit(ws_raw)

    # ── Sheets: Curated_[param] ───────────────────────────────────────────────
    n_per_sub = max(1, len(cells) // max(1, len(substrate_groups)))
    for param, (mn, mx) in thresholds.items():
        if progress_callback:
            progress_callback(f"Building Curated_{param} sheet...")
        passing = [c for c in cells if c.passes(param, mn, mx, scan=threshold_scan)]
        ws_cur  = wb.create_sheet(f"Curated_{param}")
        ws_cur.cell(1, 1, f"Threshold: {param}  [{mn} – {mx}]").font = Font(
            bold=True, size=11, color="1565C0")
        ws_cur.cell(2, 1, f"Scan filter: {'Reverse only' if threshold_scan == 'rev' else 'Both scans'}").font = Font(
            italic=True, size=10)
        ws_cur.cell(3, 1, f"Cells passing: {len(passing)} / {len(cells)}").font = Font(
            bold=True, size=10)
        row = _write_headers(ws_cur, 5)
        p_grps = group_by_substrate(passing, n_per_substrate=n_per_sub)
        for sub_name, sub_cells in p_grps.items():
            ws_cur.cell(row, FWD_DEV, f"── {sub_name} ──").font = Font(
                bold=True, italic=True, size=9, color="5D4037")
            row += 1
            for cell in sub_cells:
                _write_cell_row(ws_cur, row, cell, highlight=False)
                row += 1
            row = _write_stats_both(ws_cur, row, sub_cells)
            row += 1
        _autofit(ws_cur)

    # ── Sheet: Statistics ────────────────────────────────────────────────────
    if progress_callback:
        progress_callback("Building Statistics sheet...")
    ws_stat = wb.create_sheet("Statistics")
    ws_stat.cell(1, 1, "Batch Statistics").font = Font(bold=True, size=13, color="1565C0")
    ws_stat.cell(2, 1, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font = Font(
        italic=True, size=10)
    ws_stat.cell(3, 1, f"Total cells: {len(cells)}").font = Font(size=10)

    row = 5
    for k, h in enumerate(["Parameter", "Scan", "N", "Mean", "Std Dev", "Min", "Max", "Median"], 1):
        ws_stat.cell(row, k, h).fill = H_FILL
        ws_stat.cell(row, k).font = H_FONT
        ws_stat.cell(row, k).alignment = C_ALIGN
    row += 1

    for param in PARAMS:
        scans = ["rev"] if param == "HI" else ["fwd", "rev"]
        for scan in scans:
            vals = [c.get_val(param, scan) for c in cells if c.get_val(param, scan) != 0.0]
            if not vals:
                continue
            arr = np.array(vals)
            for k, v in enumerate([
                param,
                "N/A" if param == "HI" else scan.upper(),
                len(vals),
                round(float(np.mean(arr)),   4),
                round(float(np.std(arr)),    4),
                round(float(np.min(arr)),    4),
                round(float(np.max(arr)),    4),
                round(float(np.median(arr)), 4),
            ], 1):
                ws_stat.cell(row, k, v).font = N_FONT
                if k >= 4:
                    ws_stat.cell(row, k).number_format = "0.0000"
            row += 1

    row += 1
    ws_stat.cell(row, 1, "Threshold Summary").font = Font(bold=True, size=11)
    row += 1
    for param, (mn, mx) in thresholds.items():
        n_pass = sum(1 for c in cells if c.passes(param, mn, mx, scan=threshold_scan))
        ws_stat.cell(row, 1, param).font = N_FONT
        ws_stat.cell(row, 2, f"{mn} – {mx}").font = N_FONT
        ws_stat.cell(row, 3, f"{n_pass}/{len(cells)} passing").font = N_FONT
        row += 1

    for col in ws_stat.columns:
        mx_w = max((len(str(c.value or "")) for c in col), default=8)
        ws_stat.column_dimensions[get_column_letter(col[0].column)].width = max(mx_w + 2, 12)

    # ── Individual per-cell J-V sheets ───────────────────────────────────────
    used_sheet_names = set(wb.sheetnames)
    for i, cell in enumerate(cells):
        if progress_callback:
            progress_callback(f"Adding sheet {i+1}/{len(cells)}: {cell.device}")
        _add_individual_sheet(wb, cell, used_sheet_names)

    wb.save(output_path)
    return output_path
