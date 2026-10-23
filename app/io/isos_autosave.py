"""
CSV autosave for long-running ISOS stress experiments.

Continuous protocols (L-1, V-1 — added in later sessions) accumulate many data
points over hours or days. To survive a crash or power loss, rows are streamed
to a plain CSV file with an fsync after every append. When the experiment ends,
``consolidate_csv_into_xlsx`` folds the CSV into the tracking workbook as a new
sheet and marks the CSV as consumed.
"""

import csv
import os

import openpyxl


class CSVStream:
    """
    Append-only CSV writer with durable (fsync'd) writes.

    Usage:
        stream = CSVStream(path, ["t", "v", "j"])
        stream.open()                 # writes the header row
        stream.append({"t": 0.0, "v": 0.8, "j": 19.2})
        ...
        stream.close()
    """

    def __init__(self, path: str, columns: list):
        self.path    = path
        self.columns = list(columns)
        self._fh     = None
        self._writer = None

    def open(self):
        """Open the file and write the header row (durably)."""
        # newline="" is required for the csv module to control line endings.
        self._fh = open(self.path, "w", newline="", encoding="utf-8")
        self._writer = csv.DictWriter(self._fh, fieldnames=self.columns,
                                      extrasaction="ignore")
        self._writer.writeheader()
        self._flush()
        return self

    def append(self, row: dict):
        """Append one row and durably flush it to disk."""
        if self._writer is None:
            raise RuntimeError("CSVStream.append() called before open()")
        # Fill any missing columns with empty strings for a rectangular file.
        full = {c: row.get(c, "") for c in self.columns}
        self._writer.writerow(full)
        self._flush()

    def close(self):
        """Flush and close the underlying file handle."""
        if self._fh is not None:
            try:
                self._flush()
            finally:
                self._fh.close()
                self._fh = None
                self._writer = None

    def _flush(self):
        self._fh.flush()
        os.fsync(self._fh.fileno())

    def __enter__(self):
        return self.open()

    def __exit__(self, exc_type, exc, tb):
        self.close()
        return False


def consolidate_csv_into_xlsx(csv_path: str, xlsx_path: str, sheet_name: str) -> None:
    """
    Append the rows of ``csv_path`` as a new sheet ``sheet_name`` in the workbook
    at ``xlsx_path`` (created if it does not exist), then rename the CSV to
    ``<csv_path>.consumed`` to mark it as archived.
    """
    with open(csv_path, "r", newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))

    if os.path.exists(xlsx_path):
        wb = openpyxl.load_workbook(xlsx_path)
    else:
        wb = openpyxl.Workbook()
        # Drop the default empty sheet so we only keep what we add.
        default = wb.active
        wb.remove(default)

    ws = wb.create_sheet(sheet_name)
    for row in rows:
        ws.append(row)

    wb.save(xlsx_path)
    wb.close()

    consumed = csv_path + ".consumed"
    if os.path.exists(consumed):
        os.remove(consumed)
    os.replace(csv_path, consumed)
