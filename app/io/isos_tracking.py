"""
ISOS protocol tracking file management (generalized across D-1 / L-1 / V-1).

This module supersedes the old ``isos_exporter`` module. The generic functions
work for any ISOS protocol; the ``*_isos_d1_*`` wrappers at the bottom are thin
shims kept for backward compatibility so existing callers and old D-1 xlsx
files keep working unchanged.

Tracking file layout (openpyxl, read/write):
  ISOS_Config  — row 1: headers, row 2: experiment config + measurement params
  ISOS_Log     — row 1: headers, rows 2+: one row per measurement (cumulative)
  JV_#001_t0000h  — full MeasurementData-compatible sheet for measurement #1
  JV_#002_t0024h  — ...

Old D-1 files used the sheet names ``ISOS_D1_Config`` / ``ISOS_D1_Log``; the
readers fall back to those names so legacy files still open.

Each JV_* sheet uses the identical column layout as the standard MeasurementData
sheet exported by save_data() in exporter.py:
  A = Forward Voltage (V)         D = Reverse Voltage (V)
  B = Forward Current Density     E = Reverse Current Density
  C = Forward Power Density       F = Reverse Power Density
  Row 1 = headers (same labels)
  Row 2 = scalar params (Voc, Jsc, FF, PCE, Rs, Rsh… in cols G onward)
  Row 3+ = V/J/P arrays
  An embedded Excel scatter chart is added to the right of the data.

Individual backup files are saved separately using save_data() from exporter.py.
"""

import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import ScatterChart, Reference, Series

# Generic sheet names.
_LOG_SHEET = "ISOS_Log"
_CFG_SHEET = "ISOS_Config"

# Legacy D-1 sheet names — read-only fallback so old files still open.
_LEGACY_LOG_SHEET = "ISOS_D1_Log"
_LEGACY_CFG_SHEET = "ISOS_D1_Config"

# Original D-1 log columns (positions 1-22 must stay fixed — see _LOG_COL_* below).
_LOG_HEADERS_CORE = [
    "#", "Timestamp", "Elapsed_h", "Delta_h",
    "Voc_fwd (V)", "Jsc_fwd (mA/cm²)", "FF_fwd (%)", "PCE_fwd (%)", "PCE_fwd_norm",
    "Voc_rev (V)", "Jsc_rev (mA/cm²)", "FF_rev (%)", "PCE_rev (%)", "PCE_rev_norm",
    "HI (%)",
    "Rs_fwd (Ω·cm²)", "Rsh_fwd (Ω·cm²)",
    "Rs_rev (Ω·cm²)", "Rsh_rev (Ω·cm²)",
    "Illumination (mW/cm²)", "Scan_Rate (mV/s)", "Notes",
]

# New nullable log columns (used by sessions 4-6). Each entry maps the column
# header to the meas-dict key it is sourced from in append_isos_measurement.
_LOG_EXTRA = [
    ("Load",                       "load"),
    ("Setpoint_V",                 "setpoint_v"),
    ("MPP_Engine",                 "mpp_engine"),
    ("Stress_With_Light",          "stress_with_light"),
    ("Stress_Elapsed_h",           "stress_elapsed_h"),
    ("PCE_0_baseline_median",      "pce_0_baseline_median"),
    ("PCE_0_baseline_std",         "pce_0_baseline_std"),
    ("baseline_method",            "baseline_method"),
    ("pre_scan_relaxation_s",      "pre_scan_relaxation_s"),
    ("PCE_SPO",                    "pce_spo"),
    ("PCE_SPO_std",                "pce_spo_std"),
    ("SPO_duration_s",             "spo_duration_s"),
    ("HI_dark",                    "hi_dark"),
    ("failure_detected",           "failure_detected"),
    ("failure_reason",             "failure_reason"),
    ("R_s_dark_fwd",               "r_s_dark_fwd"),
    ("R_sh_dark_fwd",              "r_sh_dark_fwd"),
    ("R_s_dark_rev",               "r_s_dark_rev"),
    ("R_sh_dark_rev",              "r_sh_dark_rev"),
    ("n_dark_rev",                 "n_dark_rev"),
    ("J_s_dark_rev",               "j_s_dark_rev"),
    ("n_dark_rev_R2",              "n_dark_rev_r2"),
    ("Voc_setpoint",               "voc_setpoint"),
    ("Vmpp_setpoint",              "vmpp_setpoint"),
    ("I_dark_during_stress_mean",  "i_dark_during_stress_mean"),
]

_LOG_HEADERS = _LOG_HEADERS_CORE + [h for h, _ in _LOG_EXTRA]

# Column index (1-based) of key values in the Log sheet — used for normalization.
# These live in the fixed core block, so they are unaffected by the extra columns.
_LOG_COL_PCE_FWD = 8   # PCE_fwd
_LOG_COL_PCE_REV = 13  # PCE_rev

_CFG_KEYS = [
    "protocol_id",
    "sample_name", "cell_area", "detector_power", "detector_area",
    "start_voltage", "end_voltage", "step_voltage",
    "delay_ms", "nplc", "four_wire",
    "prebias_enabled", "prebias_voltage", "prebias_duration",
    "light_soaking_enabled", "light_soaking_duration",
    "dark_jv", "electrical_load",
    "scan_rate_mode", "target_scan_rate",
    "t0_timestamp",
    # new config columns (used by sessions 4-6; nullable)
    "equilibration_duration_h", "equilibration_done",
    "baseline_n_scans", "baseline_rest_s",
    "tiered_schedule_json",
]

# Column headers for each per-measurement JV sheet — matches MeasurementData exactly
# for cols A-AA, then adds ISOS-specific columns.
_JV_HEADERS = [
    "Forward Voltage (V)", "Forward Current Density (mA.cm⁻²)",
    "Forward Power Density (mW.cm⁻²)", "Reverse Voltage (V)",
    "Reverse Current Density (mA.cm⁻²)", "Reverse Power Density (mW.cm⁻²)",
    "Solar Cell Area (cm²)",
    "Forward Voc (V)", "Forward Jsc (mA.cm⁻²)",
    "Forward FF (%)", "Forward PCE (%)",
    "Forward Rs (Ω.cm²)", "Forward Rsh (Ω.cm²)",
    "Reverse Voc (V)", "Reverse Jsc (mA.cm⁻²)",
    "Reverse FF (%)", "Reverse PCE (%)",
    "Reverse Rs (Ω.cm²)", "Reverse Rsh (Ω.cm²)",
    "Hysteresis Index HI (%)",
    "Delay/point (ms)", "NPLC", "4-Wire Sensing",
    "Pre-bias Voltage (V)", "Pre-bias Duration (s)",
    "Instrument IDN", "Scan Rate (mV/s)",
    "Elapsed (h)", "Measurement #",
]

_H_FILL    = PatternFill("solid", fgColor="1565C0")
_H_FONT    = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_V_FONT    = Font(name="Calibri", size=11)
_B_FONT    = Font(name="Calibri", size=11, bold=True)
_DATA_FONT = Font(name="Calibri", size=10)
_C_ALIGN   = Alignment(horizontal="center", vertical="center")
_L_ALIGN   = Alignment(horizontal="left",   vertical="center")
_THIN      = Side(style="thin")
_H_BORDER  = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header_row(ws, headers, col_width=16):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill      = _H_FILL
        cell.font      = _H_FONT
        cell.alignment = _C_ALIGN
        cell.border    = _H_BORDER
        ws.column_dimensions[get_column_letter(c)].width = col_width


def _protocol_file_label(protocol_id: str) -> str:
    """'ISOS-D-1' → 'ISOS-D1' for use in file names (matches legacy naming)."""
    parts = str(protocol_id).split("-")
    if len(parts) <= 1:
        return str(protocol_id)
    return parts[0] + "-" + "".join(parts[1:])


def _resolve_sheet(wb, primary, legacy):
    """Return the worksheet under ``primary`` name, falling back to ``legacy``."""
    if primary in wb.sheetnames:
        return wb[primary]
    if legacy in wb.sheetnames:
        return wb[legacy]
    return None


def _read_headers(ws):
    """Read the contiguous header labels from row 1 (stops at the first blank)."""
    headers = []
    c = 1
    while True:
        h = ws.cell(1, c).value
        if h is None:
            break
        headers.append(h)
        c += 1
    return headers


# ─────────────────────────────────────────────────────────────────────────────
# Generic API
# ─────────────────────────────────────────────────────────────────────────────

def create_isos_tracking_file(folder: str, config: dict) -> str:
    """Create a new ISOS tracking workbook and return its path."""
    protocol_id = config.get("protocol_id", "ISOS-D-1") or "ISOS-D-1"
    ts    = datetime.now().strftime("%Y%m%d_%H%M%S")
    name  = config.get("sample_name", "Sample") or "Sample"
    label = _protocol_file_label(protocol_id)
    fp    = os.path.join(folder, f"{name}_{label}_{ts}.xlsx")

    wb = openpyxl.Workbook()

    # ── Config sheet ──────────────────────────────────────────────────────────
    ws_cfg        = wb.active
    ws_cfg.title  = _CFG_SHEET
    _header_row(ws_cfg, _CFG_KEYS, col_width=22)
    cfg = dict(config)
    cfg["protocol_id"] = protocol_id
    cfg["t0_timestamp"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    row = []
    for k in _CFG_KEYS:
        v = cfg.get(k, "")
        row.append(v if v is not None else "")
    ws_cfg.append(row)
    for c in range(1, len(_CFG_KEYS) + 1):
        ws_cfg.cell(2, c).font      = _V_FONT
        ws_cfg.cell(2, c).alignment = _L_ALIGN

    # ── Log sheet ─────────────────────────────────────────────────────────────
    ws_log = wb.create_sheet(_LOG_SHEET)
    _header_row(ws_log, _LOG_HEADERS, col_width=18)
    ws_log.freeze_panes = "A2"

    wb.save(fp)
    wb.close()
    return fp


def append_isos_measurement(tracking_path: str, meas: dict) -> None:
    """
    Append one measurement to the tracking file.

    Required meas keys:
      v_fwd, j_fwd, v_rev, j_rev  — array-like (numpy or list)
      Voc_fwd, Jsc_fwd, FF_fwd, PCE_fwd, Rs_fwd, Rsh_fwd
      Voc_rev, Jsc_rev, FF_rev, PCE_rev, Rs_rev, Rsh_rev
      HI, illumination, scan_rate

    Optional (used for the per-measurement JV sheet scalars):
      cell_area, delay_ms, nplc, four_wire
      prebias_enabled, prebias_voltage, prebias_duration
      instrument_idn, notes

    Optional (new nullable Log columns — see _LOG_EXTRA): load, setpoint_v,
      mpp_engine, stress_with_light, … (all default to None if absent).
    """
    wb     = openpyxl.load_workbook(tracking_path)
    ws_log = _resolve_sheet(wb, _LOG_SHEET, _LEGACY_LOG_SHEET)
    ws_cfg = _resolve_sheet(wb, _CFG_SHEET, _LEGACY_CFG_SHEET)

    # Read t0 from config sheet (scan by header name).
    t0_str = ""
    if ws_cfg is not None:
        for c in range(1, ws_cfg.max_column + 1):
            if ws_cfg.cell(1, c).value == "t0_timestamp":
                t0_str = ws_cfg.cell(2, c).value or ""
                break
    try:
        t0 = datetime.strptime(str(t0_str), "%Y-%m-%d %H:%M:%S")
    except Exception:
        t0 = datetime.now()

    now       = datetime.now()
    elapsed_h = round((now - t0).total_seconds() / 3600.0, 3)

    # Existing measurement count (rows minus header)
    n_existing = ws_log.max_row - 1

    # Delta_h and t0 normalisation reference
    if n_existing <= 0:
        delta_h      = 0.0
        pce_fwd_t0   = meas["PCE_fwd"]
        pce_rev_t0   = meas["PCE_rev"]
    else:
        last_row     = ws_log.max_row
        last_elapsed = ws_log.cell(last_row, 3).value or 0
        delta_h      = round(elapsed_h - float(last_elapsed), 3)
        pce_fwd_t0   = ws_log.cell(2, _LOG_COL_PCE_FWD).value or meas["PCE_fwd"]
        pce_rev_t0   = ws_log.cell(2, _LOG_COL_PCE_REV).value or meas["PCE_rev"]

    # Normalized PCE — ratio relative to t₀ (1.0 = same as first measurement)
    pce_fwd_norm = round(meas["PCE_fwd"] / pce_fwd_t0, 4) if pce_fwd_t0 else 1.0
    pce_rev_norm = round(meas["PCE_rev"] / pce_rev_t0, 4) if pce_rev_t0 else 1.0

    m = n_existing + 1
    log_row = [
        m,
        now.strftime("%Y-%m-%d %H:%M:%S"),
        elapsed_h,
        delta_h,
        round(float(meas["Voc_fwd"]), 4), round(float(meas["Jsc_fwd"]), 4),
        round(float(meas["FF_fwd"]),  2),  round(float(meas["PCE_fwd"]), 4),
        pce_fwd_norm,
        round(float(meas["Voc_rev"]), 4), round(float(meas["Jsc_rev"]), 4),
        round(float(meas["FF_rev"]),  2),  round(float(meas["PCE_rev"]), 4),
        pce_rev_norm,
        round(float(meas["HI"]),      2),
        round(float(meas["Rs_fwd"]),  4),  round(float(meas["Rsh_fwd"]), 4),
        round(float(meas["Rs_rev"]),  4),  round(float(meas["Rsh_rev"]), 4),
        round(float(meas.get("illumination", 0)), 2),
        round(float(meas.get("scan_rate", 0)),    2),
        meas.get("notes", ""),
    ]
    # New nullable columns — default to None when the caller does not supply them.
    log_row += [meas.get(key) for _, key in _LOG_EXTRA]
    ws_log.append(log_row)

    # Style the new data row
    r = ws_log.max_row
    for c in range(1, len(log_row) + 1):
        ws_log.cell(r, c).font      = _V_FONT
        ws_log.cell(r, c).alignment = _C_ALIGN

    # ── Per-measurement JV sheet (MeasurementData-compatible format) ──────────
    elapsed_label = f"t{int(elapsed_h):04d}h"
    sheet_num     = f"#{m:03d}"
    sheet_name    = f"JV_{sheet_num}_{elapsed_label}"

    ws_jv = wb.create_sheet(sheet_name)
    _header_row(ws_jv, _JV_HEADERS, col_width=15)

    # Adjust column widths to match MeasurementData convention
    for col_letter in ("A", "B", "C", "D", "E", "F"):
        ws_jv.column_dimensions[col_letter].width = 22
    for i in range(7, len(_JV_HEADERS) + 1):
        ws_jv.column_dimensions[get_column_letter(i)].width = 18

    # Row 2: scalar parameters in the same column positions as MeasurementData
    pb_en = meas.get("prebias_enabled", False)
    pb_v  = float(meas.get("prebias_voltage",  0)) if pb_en else 0.0
    pb_t  = float(meas.get("prebias_duration", 0)) if pb_en else 0.0

    scalar_vals = [None] * len(_JV_HEADERS)
    # cols 7-29 (1-based) ↔ indices 6-28 (0-based)
    scalar_vals[6]  = meas.get("cell_area")                           # G: cell area
    scalar_vals[7]  = round(float(meas["Voc_fwd"]), 4)               # H
    scalar_vals[8]  = round(float(meas["Jsc_fwd"]), 4)               # I
    scalar_vals[9]  = round(float(meas["FF_fwd"]),  2)               # J
    scalar_vals[10] = round(float(meas["PCE_fwd"]), 4)               # K
    scalar_vals[11] = round(float(meas["Rs_fwd"]),  4)               # L
    scalar_vals[12] = round(float(meas["Rsh_fwd"]), 4)               # M
    scalar_vals[13] = round(float(meas["Voc_rev"]), 4)               # N
    scalar_vals[14] = round(float(meas["Jsc_rev"]), 4)               # O
    scalar_vals[15] = round(float(meas["FF_rev"]),  2)               # P
    scalar_vals[16] = round(float(meas["PCE_rev"]), 4)               # Q
    scalar_vals[17] = round(float(meas["Rs_rev"]),  4)               # R
    scalar_vals[18] = round(float(meas["Rsh_rev"]), 4)               # S
    scalar_vals[19] = round(float(meas["HI"]),      2)               # T
    scalar_vals[20] = meas.get("delay_ms")                            # U
    scalar_vals[21] = meas.get("nplc")                                # V
    scalar_vals[22] = 1 if meas.get("four_wire") else 0              # W
    scalar_vals[23] = pb_v                                            # X
    scalar_vals[24] = pb_t                                            # Y
    scalar_vals[25] = meas.get("instrument_idn", "")                 # Z
    scalar_vals[26] = round(float(meas.get("scan_rate", 0)), 2)      # AA
    scalar_vals[27] = elapsed_h                                       # AB
    scalar_vals[28] = m                                               # AC

    ws_jv.append(scalar_vals)
    for c in range(1, len(_JV_HEADERS) + 1):
        ws_jv.cell(2, c).font      = _B_FONT
        ws_jv.cell(2, c).alignment = _C_ALIGN

    # V/J/P data arrays (row 3+)
    v_fwd = list(meas["v_fwd"])
    j_fwd = list(meas["j_fwd"])
    v_rev = list(meas["v_rev"])
    j_rev = list(meas["j_rev"])
    p_fwd = [float(v) * float(j) for v, j in zip(v_fwd, j_fwd)]
    p_rev = [float(v) * float(j) for v, j in zip(v_rev, j_rev)]

    nf = _DATA_FONT

    for i in range(max(len(v_fwd), len(v_rev))):
        row_data = [None] * 6
        if i < len(v_fwd):
            row_data[0] = float(v_fwd[i])
            row_data[1] = float(j_fwd[i])
            row_data[2] = float(p_fwd[i])
        if i < len(v_rev):
            row_data[3] = float(v_rev[i])
            row_data[4] = float(j_rev[i])
            row_data[5] = float(p_rev[i])
        ws_jv.append(row_data)

    for data_row in range(3, ws_jv.max_row + 1):
        for col in range(1, 7):
            cell = ws_jv.cell(data_row, col)
            cell.font           = nf
            cell.number_format  = "0.000000"
            cell.alignment      = _C_ALIGN

    # ── Embedded Excel scatter chart (J-V curve) ──────────────────────────────
    n_fwd = len(v_fwd)
    n_rev = len(v_rev)

    chart = ScatterChart()
    chart.title         = f"J-V  {sheet_num}  ({elapsed_label})"
    chart.scatterStyle  = "line"
    chart.style         = 10
    chart.x_axis.title  = "Voltage (V)"
    chart.y_axis.title  = "J (mA/cm²)"
    chart.width         = 18   # cm
    chart.height        = 12   # cm
    chart.legend.position = "b"

    # Forward series (cols A, B)
    x_fwd = Reference(ws_jv, min_col=1, min_row=3, max_row=2 + n_fwd)
    y_fwd = Reference(ws_jv, min_col=2, min_row=3, max_row=2 + n_fwd)
    s_fwd = Series(y_fwd, x_fwd, title="Forward")
    s_fwd.graphicalProperties.line.solidFill = "1565C0"
    s_fwd.graphicalProperties.line.width = 19050  # 1.5 pt in EMU
    chart.series.append(s_fwd)

    # Reverse series (cols D, E)
    x_rev = Reference(ws_jv, min_col=4, min_row=3, max_row=2 + n_rev)
    y_rev = Reference(ws_jv, min_col=5, min_row=3, max_row=2 + n_rev)
    s_rev = Series(y_rev, x_rev, title="Reverse")
    s_rev.graphicalProperties.line.solidFill = "B71C1C"
    s_rev.graphicalProperties.line.width = 19050
    chart.series.append(s_rev)

    # Place chart to the right of the scalar data area (col AC + a gap → col AF)
    ws_jv.add_chart(chart, "AF1")

    wb.save(tracking_path)
    wb.close()


def read_isos_config(tracking_path: str) -> dict:
    """Read config values from the tracking file. Returns {} on failure."""
    try:
        wb = openpyxl.load_workbook(tracking_path, data_only=True)
        ws = _resolve_sheet(wb, _CFG_SHEET, _LEGACY_CFG_SHEET)
        if ws is None:
            wb.close()
            return {}
        out = {}
        for c in range(1, ws.max_column + 1):
            key = ws.cell(1, c).value
            val = ws.cell(2, c).value
            if key:
                out[key] = val
        wb.close()
        return out
    except Exception:
        return {}


def read_isos_log(tracking_path: str) -> list:
    """Read all measurement rows from the log. Returns list of dicts."""
    try:
        wb = openpyxl.load_workbook(tracking_path, data_only=True)
        ws = _resolve_sheet(wb, _LOG_SHEET, _LEGACY_LOG_SHEET)
        if ws is None:
            wb.close()
            return []
        headers = _read_headers(ws)
        rows    = []
        for r in range(2, ws.max_row + 1):
            row = {headers[c - 1]: ws.cell(r, c).value
                   for c in range(1, len(headers) + 1)}
            if row.get("#") is not None:
                rows.append(row)
        wb.close()
        return rows
    except Exception:
        return []


def read_isos_jv_curves(tracking_path: str) -> list:
    """
    Read all JV curve sheets from the tracking file.
    Returns list of dicts: {direction, elapsed_h, v, j}
    where direction is 'Fwd' or 'Rev'.

    Sheet name format: JV_#NNN_tXXXXh
    Columns: A/B = V_fwd/J_fwd, D/E = V_rev/J_rev (same as MeasurementData)
    Data starts at row 3 (row 1 = headers, row 2 = scalar params).
    """
    try:
        wb     = openpyxl.load_workbook(tracking_path, data_only=True)
        curves = []
        for sheet_name in wb.sheetnames:
            if not sheet_name.startswith("JV_"):
                continue
            parts = sheet_name.split("_")
            # New format: JV_#NNN_tXXXXh  (3 parts)
            if len(parts) < 3:
                continue
            elapsed_label = parts[2]  # e.g. t0024h
            try:
                elapsed_h = int(elapsed_label[1:-1])  # strip 't' and 'h'
            except Exception:
                elapsed_h = 0
            ws = wb[sheet_name]
            v_fwd, j_fwd = [], []
            v_rev, j_rev = [], []
            # Data starts at row 3 in the new format
            start_row = 3
            for r in range(start_row, ws.max_row + 1):
                vf = ws.cell(r, 1).value
                jf = ws.cell(r, 2).value
                if vf is not None and jf is not None:
                    v_fwd.append(float(vf))
                    j_fwd.append(float(jf))
                vr = ws.cell(r, 4).value
                jr = ws.cell(r, 5).value
                if vr is not None and jr is not None:
                    v_rev.append(float(vr))
                    j_rev.append(float(jr))
            if v_fwd:
                curves.append({
                    "direction": "Fwd",
                    "elapsed_h": elapsed_h,
                    "v": v_fwd,
                    "j": j_fwd,
                })
            if v_rev:
                curves.append({
                    "direction": "Rev",
                    "elapsed_h": elapsed_h,
                    "v": v_rev,
                    "j": j_rev,
                })
        wb.close()
        return curves
    except Exception:
        return []


# ─────────────────────────────────────────────────────────────────────────────
# Generic auxiliary sheets (baseline scans, SPO traces, equilibration logs)
#
# Continuous protocols (L-1 / V-1) persist several supporting tables alongside
# the main Log: baseline conditioning scans, SPO traces after each scan, and an
# equilibration log. These two helpers write those without imposing a fixed
# schema, so the runner owns the column definitions.
# ─────────────────────────────────────────────────────────────────────────────

def append_curve_sheet(tracking_path: str, sheet_name: str, columns: dict) -> None:
    """Write ``columns`` (an ordered ``{header: list}`` mapping) as a fresh sheet.

    Used for Baseline_#NN and SPO_#NNN_tXXXXh sheets. Any pre-existing sheet of
    the same name is replaced so re-runs stay idempotent. Columns may differ in
    length; shorter columns are padded with blanks.
    """
    wb = openpyxl.load_workbook(tracking_path)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    headers = list(columns.keys())
    _header_row(ws, headers, col_width=18)

    n_rows = max((len(v) for v in columns.values()), default=0)
    for i in range(n_rows):
        row = []
        for h in headers:
            col = columns[h]
            row.append(col[i] if i < len(col) else None)
        ws.append(row)

    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            ws.cell(r, c).font      = _DATA_FONT
            ws.cell(r, c).alignment = _C_ALIGN

    wb.save(tracking_path)
    wb.close()


def append_log_sheet_row(tracking_path: str, sheet_name: str,
                         headers: list, row: list) -> None:
    """Append ``row`` to a named log sheet, creating it with ``headers`` if new.

    Used for the Equilibration_log sheet. Unlike ``append_curve_sheet`` this is
    additive — each call appends one row to the existing table.
    """
    wb = openpyxl.load_workbook(tracking_path)
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)
        _header_row(ws, headers, col_width=20)

    ws.append(list(row))
    r = ws.max_row
    for c in range(1, len(row) + 1):
        ws.cell(r, c).font      = _V_FONT
        ws.cell(r, c).alignment = _C_ALIGN

    wb.save(tracking_path)
    wb.close()


# ─────────────────────────────────────────────────────────────────────────────
# Backward-compatibility wrappers (ISOS-D-1)
#
# These keep the old public names working. They forward to the generic
# functions with protocol_id="ISOS-D-1".
# ─────────────────────────────────────────────────────────────────────────────

def create_isos_d1_tracking_file(folder: str, config: dict) -> str:
    """Create a new ISOS-D-1 tracking workbook and return its path."""
    cfg = dict(config)
    cfg.setdefault("protocol_id", "ISOS-D-1")
    return create_isos_tracking_file(folder, cfg)


def append_isos_d1_measurement(tracking_path: str, meas: dict) -> None:
    """Append one ISOS-D-1 measurement to the tracking file."""
    return append_isos_measurement(tracking_path, meas)


def read_isos_d1_config(tracking_path: str) -> dict:
    """Read ISOS-D-1 config values from the tracking file."""
    return read_isos_config(tracking_path)


def read_isos_d1_log(tracking_path: str) -> list:
    """Read all ISOS-D-1 measurement rows from the log."""
    return read_isos_log(tracking_path)


def read_isos_d1_jv_curves(tracking_path: str) -> list:
    """Read all ISOS-D-1 JV curve sheets from the tracking file."""
    return read_isos_jv_curves(tracking_path)
