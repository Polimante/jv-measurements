"""
Generate a TEST_ISOS_D1_*.xlsx tracking file for manual testing.

Usage:
    python create_test_isos_file.py

Creates the file in the project root directory.
Simulates 4 measurements at t=0, 24, 48, 72 h with slight PCE degradation.
"""

import sys
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).parent))

from app.io.isos_tracking import (
    create_isos_d1_tracking_file,
    append_isos_d1_measurement,
)
from app.measurement.analysis import (
    calculate_jv_parameters,
    calculate_resistances,
    calculate_hysteresis_index,
)
import app.io.isos_tracking as _mod

# ── Experiment configuration ──────────────────────────────────────────────────

CONFIG = {
    "sample_name":            "TEST_ISOS",
    "cell_area":               0.07,
    "detector_power":          110.0,
    "detector_area":           1.13,
    "start_voltage":          -0.1,
    "end_voltage":             1.1,
    "step_voltage":            0.025,
    "delay_ms":                100.0,
    "nplc":                    1.0,
    "four_wire":               False,
    "prebias_enabled":         True,
    "prebias_voltage":         1.2,
    "prebias_duration":        2.0,
    "light_soaking_enabled":   False,
    "light_soaking_duration":  300.0,
    "dark_jv":                 False,
    "electrical_load":         "OC",
    "scan_rate_mode":          "advanced",
    "target_scan_rate":        100.0,
}

IDN = "KEITHLEY INSTRUMENTS INC.,MODEL 2400,1234567,C32   Mar  4 2004 17:45:26/A02  /K/J"
PIN = CONFIG["detector_power"] / CONFIG["detector_area"]  # ~97.3 mW/cm²

# Measurements: (elapsed_hours, degradation_factor)
MEASUREMENTS = [
    (0,   1.000),
    (24,  0.975),
    (48,  0.950),
    (72,  0.920),
]

# ── Synthetic J-V curve ───────────────────────────────────────────────────────

def make_jv(jsc, voc_f, voc_r, rsh=800.0):
    """Return (v_fwd, j_fwd, v_rev, j_rev) as numpy arrays with realistic hysteresis."""
    v = np.linspace(-0.1, 1.1, 49)
    j0 = 1e-10
    kT = 0.026

    def j_curve(voltage, voc_target, rsh_val):
        # Scale J0 so that J(voc_target) ≈ 0
        _j0 = jsc / (np.exp(voc_target / kT) - 1)
        # Clip at 500 (well above any realistic voc/kT) to avoid float64 overflow
        j = jsc - _j0 * (np.exp(np.clip(voltage / kT, -300, 500)) - 1) - voltage / rsh_val
        return np.clip(j, -jsc * 0.15, jsc * 1.02)

    j_f = j_curve(v,        voc_f, rsh_val=rsh)
    j_r = j_curve(v[::-1],  voc_r, rsh_val=rsh * 1.05)
    return v, j_f, v[::-1], j_r


# ── Fake datetime subclass so elapsed times land at round values ──────────────

_real_datetime = datetime


class _FakeDatetime(_real_datetime):
    """datetime subclass that overrides now() with a controlled fake time."""
    _fake_now = None

    @classmethod
    def now(cls, tz=None):
        return cls._fake_now

    @classmethod
    def strptime(cls, date_string, format_str):
        return _real_datetime.strptime(date_string, format_str)


# ── Create tracking file ──────────────────────────────────────────────────────

output_dir = str(Path(__file__).parent)
path = create_isos_d1_tracking_file(output_dir, CONFIG)
print(f"Created tracking file:\n  {path}\n")

# Pin t0 so measurement #1 lands at elapsed=0
t0 = _real_datetime.now() - timedelta(seconds=1)

# Patch t0_timestamp in the config sheet
import openpyxl as _openpyxl

def _patch_t0(file_path, t0_str):
    wb = _openpyxl.load_workbook(file_path)
    ws = wb["ISOS_D1_Config"]
    for c in range(1, ws.max_column + 2):
        if ws.cell(1, c).value == "t0_timestamp":
            ws.cell(2, c).value = t0_str
            break
    wb.save(file_path)
    wb.close()


_patch_t0(path, t0.strftime("%Y-%m-%d %H:%M:%S"))

# Replace datetime in the exporter module with our controllable subclass
_mod.datetime = _FakeDatetime

nplc_ms     = CONFIG["nplc"] * (1000.0 / 60.0)
t_per_point = (CONFIG["delay_ms"] + nplc_ms + 50.0) / 1000.0
scan_rate   = (abs(CONFIG["step_voltage"]) * 1000.0) / t_per_point

for i, (elapsed_h, deg) in enumerate(MEASUREMENTS):
    _FakeDatetime._fake_now = t0 + timedelta(hours=elapsed_h)

    jsc   = 21.0 * deg
    voc_f = 1.03 * deg
    voc_r = 1.06 * deg   # higher Voc on reverse → positive HI

    v_f, j_f, v_r, j_r = make_jv(jsc, voc_f, voc_r)

    Voc_f, Jsc_f, FF_f, PCE_f = calculate_jv_parameters(v_f, j_f, PIN)
    Voc_r, Jsc_r, FF_r, PCE_r = calculate_jv_parameters(v_r, j_r, PIN)
    Rs_f,  Rsh_f               = calculate_resistances(v_f, j_f, Voc_f)
    Rs_r,  Rsh_r               = calculate_resistances(v_r, j_r, Voc_r)
    hi                          = calculate_hysteresis_index(PCE_f, PCE_r)

    append_isos_d1_measurement(path, {
        "v_fwd":   v_f,    "j_fwd":   j_f,
        "v_rev":   v_r,    "j_rev":   j_r,
        "Voc_fwd": Voc_f,  "Jsc_fwd": Jsc_f,
        "FF_fwd":  FF_f,   "PCE_fwd": PCE_f,
        "Rs_fwd":  Rs_f,   "Rsh_fwd": Rsh_f,
        "Voc_rev": Voc_r,  "Jsc_rev": Jsc_r,
        "FF_rev":  FF_r,   "PCE_rev": PCE_r,
        "Rs_rev":  Rs_r,   "Rsh_rev": Rsh_r,
        "HI":               hi,
        "illumination":     PIN,
        "scan_rate":        scan_rate,
        "cell_area":        CONFIG["cell_area"],
        "delay_ms":         CONFIG["delay_ms"],
        "nplc":             CONFIG["nplc"],
        "four_wire":        CONFIG["four_wire"],
        "prebias_enabled":  CONFIG["prebias_enabled"],
        "prebias_voltage":  CONFIG["prebias_voltage"],
        "prebias_duration": CONFIG["prebias_duration"],
        "instrument_idn":   IDN,
        "notes":            f"Synthetic — t={elapsed_h} h",
    })

    print(f"  #{i+1}  t={elapsed_h:>3} h  "
          f"PCE_fwd={PCE_f:.2f}%  PCE_rev={PCE_r:.2f}%  HI={hi:.1f}%")

# Restore real datetime
_mod.datetime = _real_datetime

print(f"\nDone.")
print("Open in ISOS Protocols panel (Alt+F8) -> 'Load Experiment...'")
print("All config parameters should populate automatically.")
