"""
Diagnostic script — reads an ISOS-D-1 tracking file and prints
exactly what is stored in its config sheet.

Usage:
    python diagnose_isos_file.py                 (opens file picker)
    python diagnose_isos_file.py path/to/file.xlsx
"""

import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl is not installed. Run: pip install openpyxl")
    sys.exit(1)

_CFG_SHEET = "ISOS_D1_Config"


def diagnose(file_path: str):
    print(f"\nFile: {file_path}")
    print("=" * 70)

    wb = openpyxl.load_workbook(file_path, data_only=True)
    print(f"Sheets found: {wb.sheetnames}\n")

    if _CFG_SHEET not in wb.sheetnames:
        print(f"ERROR: Sheet '{_CFG_SHEET}' NOT FOUND in this file.")
        print("This file is not a valid ISOS-D-1 tracking file, or was created")
        print("by an incompatible version of the software.")
        wb.close()
        return

    ws = wb[_CFG_SHEET]
    print(f"Config sheet dimensions: {ws.dimensions}")
    print(f"Max row: {ws.max_row}, Max column: {ws.max_column}\n")

    print(f"{'Column':<8} {'Header (row 1)':<30} {'Value (row 2)'}")
    print("-" * 70)
    for c in range(1, ws.max_column + 2):
        key = ws.cell(1, c).value
        val = ws.cell(2, c).value
        if key is None and val is None:
            break
        print(f"  {c:<6} {str(key):<30} {repr(val)}")

    wb.close()
    print("\n" + "=" * 70)
    print("If all 'Value (row 2)' entries show None, the config data")
    print("was never written to this file (likely a file creation bug).")
    print("If the headers differ from the expected keys, the file was")
    print("created with a different version of the software.")


if __name__ == "__main__":
    if len(sys.argv) > 1:
        path = sys.argv[1]
    else:
        try:
            import tkinter as tk
            from tkinter import filedialog
            root = tk.Tk()
            root.withdraw()
            path = filedialog.askopenfilename(
                title="Select ISOS-D-1 Tracking File",
                filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            )
            root.destroy()
        except Exception:
            path = input("Enter path to the ISOS-D-1 tracking file: ").strip()

    if not path:
        print("No file selected.")
        sys.exit(0)

    try:
        diagnose(path)
    except Exception as e:
        print(f"ERROR reading file: {e}")

    input("\nPress Enter to exit...")
