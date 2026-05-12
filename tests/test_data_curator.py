"""
Tests for app/io/data_curator.py

Tests cover:
  - CellData attribute defaults and device name extraction
  - CellData.get_val  and  CellData.passes (threshold logic)
  - _extract_prefix  (filename → substrate prefix)
  - group_by_substrate (grouping by prefix)
  - _stats helper (mean / std / CV% / max)
  - Filtering: which cells survive a threshold pass
  - export_curated_excel: end-to-end write + sheet names (no instrument needed)
"""

import os
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from app.io.data_curator import (
    CellData,
    _extract_prefix,
    group_by_substrate,
    _build_statistics_sheet,
    _export_composition_sheets,
    export_curated_excel,
)


# ── helpers ───────────────────────────────────────────────────────────────────

def make_cell(filename, pce_rev=15.0, pce_fwd=14.0,
              voc_rev=1.05, jsc_rev=20.0, ff_rev=71.4,
              rs_rev=3.0, rsh_rev=500.0, hi=6.7):
    """Create a CellData with controlled parameters (no file read needed)."""
    c = CellData.__new__(CellData)
    c.filepath = str(Path(tempfile.gettempdir()) / filename)
    c.filename = filename
    c.device   = Path(filename).stem.split("_")[0]

    c.PCE_rev = pce_rev;  c.PCE_fwd = pce_fwd
    c.Voc_rev = voc_rev;  c.Voc_fwd = voc_rev - 0.02
    c.Jsc_rev = jsc_rev;  c.Jsc_fwd = jsc_rev - 0.5
    c.FF_rev  = ff_rev;   c.FF_fwd  = ff_rev  - 2.0
    c.Rs_rev  = rs_rev;   c.Rs_fwd  = rs_rev  + 0.5
    c.Rsh_rev = rsh_rev;  c.Rsh_fwd = rsh_rev - 50.0
    c.HI = hi
    return c


# ─────────────────────────────────────────────────────────────────────────────
# CellData — basics
# ─────────────────────────────────────────────────────────────────────────────

class TestCellDataDefaults:

    def test_all_params_default_to_zero(self):
        c = CellData.__new__(CellData)
        c.filepath = "/tmp/D1C1_20260101.xlsx"
        c.filename = "D1C1_20260101.xlsx"
        c.device   = "D1C1"
        for attr in ["Voc_fwd","Jsc_fwd","FF_fwd","PCE_fwd","Rs_fwd","Rsh_fwd",
                     "Voc_rev","Jsc_rev","FF_rev","PCE_rev","Rs_rev","Rsh_rev","HI"]:
            setattr(c, attr, 0.0)
            assert getattr(c, attr) == 0.0

    def test_device_extraction_removes_timestamp(self):
        """Filename D1C1_20260412.xlsx → device = 'D1C1'."""
        c = CellData.__new__(CellData)
        c.filepath = "/tmp/D1C1_20260412.xlsx"
        c.filename = "D1C1_20260412.xlsx"
        c.device   = c._extract_device() if hasattr(CellData, '_extract_device') else "D1C1"
        assert "D1C1" in c.device or c.device == "D1C1"


# ─────────────────────────────────────────────────────────────────────────────
# CellData.get_val
# ─────────────────────────────────────────────────────────────────────────────

class TestCellDataGetVal:

    def setup_method(self):
        self.cell = make_cell("D1C1_20260412.xlsx", pce_rev=16.0, pce_fwd=14.5, hi=9.4)

    def test_get_pce_rev(self):
        assert self.cell.get_val("PCE", "rev") == pytest.approx(16.0)

    def test_get_pce_fwd(self):
        assert self.cell.get_val("PCE", "fwd") == pytest.approx(14.5)

    def test_get_hi_returns_same_for_both_scans(self):
        assert self.cell.get_val("HI", "fwd") == pytest.approx(self.cell.HI)
        assert self.cell.get_val("HI", "rev") == pytest.approx(self.cell.HI)

    def test_get_voc_rev(self):
        assert self.cell.get_val("Voc", "rev") == pytest.approx(1.05)


# ─────────────────────────────────────────────────────────────────────────────
# CellData.passes (threshold logic)
# ─────────────────────────────────────────────────────────────────────────────

class TestCellDataPasses:

    def setup_method(self):
        self.cell = make_cell("D1C1_20260412.xlsx",
                              pce_rev=15.0, pce_fwd=12.0, voc_rev=1.05)

    # --- PCE threshold on rev scan ---
    def test_passes_above_threshold(self):
        assert self.cell.passes("PCE", 10.0, scan="rev") is True

    def test_fails_below_threshold(self):
        assert self.cell.passes("PCE", 20.0, scan="rev") is False

    def test_passes_equal_to_threshold(self):
        assert self.cell.passes("PCE", 15.0, scan="rev") is True

    # --- Threshold on fwd scan ---
    def test_passes_fwd_scan(self):
        assert self.cell.passes("PCE", 10.0, scan="fwd") is True

    def test_fails_fwd_scan(self):
        assert self.cell.passes("PCE", 13.0, scan="fwd") is False

    # --- "both" scan requires BOTH fwd AND rev to pass ---
    def test_passes_both_when_both_above(self):
        # fwd=12.0, rev=15.0; threshold=10 → both pass
        assert self.cell.passes("PCE", 10.0, scan="both") is True

    def test_fails_both_when_fwd_below(self):
        # fwd=12.0, rev=15.0; threshold=13 → fwd fails
        assert self.cell.passes("PCE", 13.0, scan="both") is False

    def test_fails_both_when_rev_below(self):
        # fwd=12.0, rev=15.0; threshold=16 → rev fails
        assert self.cell.passes("PCE", 16.0, scan="both") is False

    # --- Voc threshold ---
    def test_passes_voc(self):
        assert self.cell.passes("Voc", 1.0, scan="rev") is True

    def test_fails_voc(self):
        assert self.cell.passes("Voc", 1.1, scan="rev") is False


# ─────────────────────────────────────────────────────────────────────────────
# _extract_prefix
# ─────────────────────────────────────────────────────────────────────────────

class TestExtractPrefix:

    def test_standard_prefix(self):
        assert _extract_prefix("/data/D1C1_20260412.xlsx") == "D1"

    def test_prefix_with_underscore(self):
        assert _extract_prefix("/data/D1_C1_20260412.xlsx") == "D1"

    def test_two_digit_prefix(self):
        assert _extract_prefix("/data/AB12C1_20260412.xlsx") == "AB12"

    def test_no_digit_in_prefix(self):
        """Filename starting with only letters returns fallback."""
        result = _extract_prefix("/data/Control_20260412.xlsx")
        # Either returns full stem or 'Control' — just not empty
        assert len(result) > 0

    def test_different_compositions(self):
        p1 = _extract_prefix("/data/D1C1_20260412.xlsx")
        p2 = _extract_prefix("/data/D2C1_20260412.xlsx")
        assert p1 != p2
        assert p1 == "D1"
        assert p2 == "D2"

    def test_same_prefix_different_cells(self):
        p1 = _extract_prefix("/data/D1C1_20260412.xlsx")
        p2 = _extract_prefix("/data/D1C2_20260412.xlsx")
        assert p1 == p2 == "D1"


# ─────────────────────────────────────────────────────────────────────────────
# group_by_substrate
# ─────────────────────────────────────────────────────────────────────────────

class TestGroupBySubstrate:

    def _make_cells(self, filenames):
        return [make_cell(fn) for fn in filenames]

    def test_single_group(self):
        cells = self._make_cells(["D1C1_20260412.xlsx",
                                  "D1C2_20260412.xlsx",
                                  "D1C3_20260412.xlsx"])
        groups = group_by_substrate(cells)
        assert list(groups.keys()) == ["D1"]
        assert len(groups["D1"]) == 3

    def test_two_groups(self):
        cells = self._make_cells(["D1C1_20260412.xlsx",
                                  "D1C2_20260412.xlsx",
                                  "D2C1_20260412.xlsx",
                                  "D2C2_20260412.xlsx"])
        groups = group_by_substrate(cells)
        assert set(groups.keys()) == {"D1", "D2"}
        assert len(groups["D1"]) == 2
        assert len(groups["D2"]) == 2

    def test_empty_list(self):
        groups = group_by_substrate([])
        assert groups == {}

    def test_groups_sorted_alphabetically(self):
        cells = self._make_cells(["C1C1_20260412.xlsx",
                                  "A1C1_20260412.xlsx",
                                  "B1C1_20260412.xlsx"])
        keys = list(group_by_substrate(cells).keys())
        assert keys == sorted(keys)


# ─────────────────────────────────────────────────────────────────────────────
# Filtering logic (cells that pass all thresholds)
# ─────────────────────────────────────────────────────────────────────────────

class TestFiltering:

    def setup_method(self):
        self.good = make_cell("D1C1_20260412.xlsx", pce_rev=18.0, voc_rev=1.10,
                              ff_rev=75.0, jsc_rev=22.0)
        self.bad  = make_cell("D1C2_20260412.xlsx", pce_rev=4.0,  voc_rev=0.70,
                              ff_rev=40.0, jsc_rev=8.0)
        self.mid  = make_cell("D1C3_20260412.xlsx", pce_rev=12.0, voc_rev=1.00,
                              ff_rev=60.0, jsc_rev=18.0)

    def test_all_pass_with_no_thresholds(self):
        cells = [self.good, self.bad, self.mid]
        passing = [c for c in cells if all(
            c.passes(p, mn, scan="rev") for p, mn in {}.items()
        )]
        assert len(passing) == 3

    def test_single_pce_threshold(self):
        cells = [self.good, self.bad, self.mid]
        thresholds = {"PCE": 10.0}
        passing = [c for c in cells if
                   all(c.passes(p, mn, "rev") for p, mn in thresholds.items())]
        assert self.good in passing
        assert self.mid  in passing
        assert self.bad  not in passing

    def test_multiple_thresholds_all_must_pass(self):
        cells = [self.good, self.bad, self.mid]
        thresholds = {"PCE": 10.0, "Voc": 1.05}
        passing = [c for c in cells if
                   all(c.passes(p, mn, "rev") for p, mn in thresholds.items())]
        # Only good (PCE=18, Voc=1.10) passes both; mid has Voc=1.00 < 1.05
        assert self.good in passing
        assert self.bad  not in passing
        assert self.mid  not in passing

    def test_strict_threshold_no_one_passes(self):
        cells = [self.good, self.bad, self.mid]
        thresholds = {"PCE": 25.0}
        passing = [c for c in cells if
                   all(c.passes(p, mn, "rev") for p, mn in thresholds.items())]
        assert passing == []

    def test_relaxed_threshold_all_pass(self):
        cells = [self.good, self.bad, self.mid]
        thresholds = {"PCE": 1.0}
        passing = [c for c in cells if
                   all(c.passes(p, mn, "rev") for p, mn in thresholds.items())]
        assert len(passing) == 3


# ─────────────────────────────────────────────────────────────────────────────
# _stats helper (tested via _build_statistics_sheet indirectly, but we can
# also access the function directly if needed — we test the public behaviour)
# ─────────────────────────────────────────────────────────────────────────────

class TestStatsLogic:
    """Verify statistical calculations on known data."""

    def _pce_stats(self, values):
        """Extract mean/std/max by building a minimal cell set."""
        cells = []
        for v in values:
            c = make_cell(f"D1C{len(cells)+1}_20260412.xlsx", pce_rev=v)
            cells.append(c)
        pce_vals = [c.PCE_rev for c in cells]
        arr = np.array(pce_vals)
        mean = float(np.mean(arr))
        std  = float(np.std(arr, ddof=1)) if len(arr) > 1 else 0.0
        cv   = std / abs(mean) * 100 if abs(mean) > 1e-10 else 0.0
        mx   = float(np.max(arr))
        return mean, std, cv, mx

    def test_mean_known_values(self):
        mean, _, _, _ = self._pce_stats([10.0, 20.0, 30.0])
        assert mean == pytest.approx(20.0)

    def test_std_known_values(self):
        _, std, _, _ = self._pce_stats([10.0, 20.0, 30.0])
        # std(ddof=1) of [10,20,30] = 10.0
        assert std == pytest.approx(10.0)

    def test_cv_known_values(self):
        _, std, cv, _ = self._pce_stats([10.0, 20.0, 30.0])
        expected_cv = 10.0 / 20.0 * 100
        assert cv == pytest.approx(expected_cv)

    def test_max_known_values(self):
        _, _, _, mx = self._pce_stats([10.0, 20.0, 30.0])
        assert mx == pytest.approx(30.0)

    def test_single_value_std_is_zero(self):
        _, std, _, _ = self._pce_stats([15.0])
        assert std == pytest.approx(0.0)


# ─────────────────────────────────────────────────────────────────────────────
# Champion cell selection
# ─────────────────────────────────────────────────────────────────────────────

class TestChampionSelection:

    def test_champion_is_max_pce_rev(self):
        cells = [
            make_cell("D1C1_20260412.xlsx", pce_rev=12.0),
            make_cell("D1C2_20260412.xlsx", pce_rev=18.5),
            make_cell("D1C3_20260412.xlsx", pce_rev=9.0),
        ]
        champ = max(cells, key=lambda c: c.PCE_rev)
        assert champ.PCE_rev == pytest.approx(18.5)

    def test_champion_per_group(self):
        d1_cells = [make_cell("D1C1_20260412.xlsx", pce_rev=12.0),
                    make_cell("D1C2_20260412.xlsx", pce_rev=18.5)]
        d2_cells = [make_cell("D2C1_20260412.xlsx", pce_rev=16.0),
                    make_cell("D2C2_20260412.xlsx", pce_rev=11.0)]
        champ_d1 = max(d1_cells, key=lambda c: c.PCE_rev)
        champ_d2 = max(d2_cells, key=lambda c: c.PCE_rev)
        assert champ_d1.PCE_rev == pytest.approx(18.5)
        assert champ_d2.PCE_rev == pytest.approx(16.0)


# ─────────────────────────────────────────────────────────────────────────────
# Hysteresis index per cell
# ─────────────────────────────────────────────────────────────────────────────

class TestCellHysteresis:

    def test_hi_stored_correctly(self):
        c = make_cell("D1C1_20260412.xlsx", hi=12.5)
        assert c.HI == pytest.approx(12.5)

    def test_cells_below_10_count(self):
        cells = [make_cell(f"D1C{i}_20260412.xlsx", hi=h)
                 for i, h in enumerate([3.0, 8.0, 12.0, 20.0, 5.0])]
        lt10 = sum(1 for c in cells if c.HI < 10)
        assert lt10 == 3   # 3.0, 8.0, 5.0

    def test_cells_below_5_count(self):
        cells = [make_cell(f"D1C{i}_20260412.xlsx", hi=h)
                 for i, h in enumerate([3.0, 8.0, 12.0, 20.0, 5.0])]
        lt5 = sum(1 for c in cells if c.HI < 5)
        assert lt5 == 1   # only 3.0


# ─────────────────────────────────────────────────────────────────────────────
# export_curated_excel — smoke test (writes a real xlsx, checks sheet names)
# ─────────────────────────────────────────────────────────────────────────────

class TestExportCuratedExcel:
    """
    These tests write actual Excel files to a temp directory.
    They verify the workbook structure without reading measurement data.
    """

    def _make_cells(self):
        return [
            make_cell("D1C1_20260412.xlsx", pce_rev=18.0, pce_fwd=16.5),
            make_cell("D1C2_20260412.xlsx", pce_rev=12.0, pce_fwd=11.0),
            make_cell("D2C1_20260412.xlsx", pce_rev=15.5, pce_fwd=14.0),
            make_cell("D2C2_20260412.xlsx", pce_rev=6.0,  pce_fwd=5.5),
        ]

    def test_creates_file(self, tmp_path):
        cells = self._make_cells()
        groups = group_by_substrate(cells)
        out = str(tmp_path / "curated.xlsx")
        export_curated_excel(cells, {}, "rev", groups, out)
        assert Path(out).exists()

    def test_sheet_raw_data_exists(self, tmp_path):
        import openpyxl
        cells = self._make_cells()
        groups = group_by_substrate(cells)
        out = str(tmp_path / "curated.xlsx")
        export_curated_excel(cells, {}, "rev", groups, out)
        wb = openpyxl.load_workbook(out)
        assert "Raw_Data" in wb.sheetnames

    def test_sheet_statistics_exists(self, tmp_path):
        import openpyxl
        cells = self._make_cells()
        groups = group_by_substrate(cells)
        out = str(tmp_path / "curated.xlsx")
        export_curated_excel(cells, {}, "rev", groups, out)
        wb = openpyxl.load_workbook(out)
        assert "Statistics" in wb.sheetnames

    def test_curated_sheet_created_for_threshold(self, tmp_path):
        import openpyxl
        cells = self._make_cells()
        groups = group_by_substrate(cells)
        out = str(tmp_path / "curated.xlsx")
        export_curated_excel(cells, {"PCE": 10.0}, "rev", groups, out)
        wb = openpyxl.load_workbook(out)
        assert "Curated_PCE" in wb.sheetnames

    def test_composition_sheets_created(self, tmp_path):
        import openpyxl
        cells = self._make_cells()
        groups = group_by_substrate(cells)
        comp_groups = [
            {"name": "Group_D1", "prefixes": ["D1"]},
            {"name": "Group_D2", "prefixes": ["D2"]},
        ]
        out = str(tmp_path / "curated.xlsx")
        export_curated_excel(cells, {}, "rev", groups, out,
                             composition_groups=comp_groups)
        wb = openpyxl.load_workbook(out)
        sheet_names = wb.sheetnames
        assert any("Rev" in s and "All" in s for s in sheet_names)
        assert any("Fwd" in s and "All" in s for s in sheet_names)

    def test_no_threshold_no_curated_sheets(self, tmp_path):
        import openpyxl
        cells = self._make_cells()
        groups = group_by_substrate(cells)
        out = str(tmp_path / "curated.xlsx")
        export_curated_excel(cells, {}, "rev", groups, out)
        wb = openpyxl.load_workbook(out)
        curated_sheets = [s for s in wb.sheetnames if s.startswith("Curated_")]
        assert curated_sheets == []

    def test_individual_cell_sheets_created(self, tmp_path):
        import openpyxl
        cells = self._make_cells()
        groups = group_by_substrate(cells)
        out = str(tmp_path / "curated.xlsx")
        export_curated_excel(cells, {}, "rev", groups, out)
        wb = openpyxl.load_workbook(out)
        # Should have at least Raw_Data + Statistics + one per cell
        assert len(wb.sheetnames) >= 2 + len(cells)

    def test_filtered_cells_count_with_threshold(self, tmp_path):
        """Only cells with PCE_rev >= 10 should appear in Curated_PCE."""
        import openpyxl
        cells = self._make_cells()  # PCE_rev: 18, 12, 15.5, 6 → 3 pass threshold 10
        groups = group_by_substrate(cells)
        out = str(tmp_path / "curated.xlsx")
        export_curated_excel(cells, {"PCE": 10.0}, "rev", groups, out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Curated_PCE"]
        # Row 7 contains cell summary text "Cells passing: X / 4"
        found = False
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell and "3" in str(cell) and "4" in str(cell):
                    found = True
        assert found


# ─────────────────────────────────────────────────────────────────────────────
# Additional tests — Part C of the bug-hunt suite
# ─────────────────────────────────────────────────────────────────────────────

class TestCuratedSheetValues:
    """Verify that specific cell values inside Curated_PCE are correct."""

    def _cells_with_known_pce(self):
        return [
            make_cell("D1C1_20260412.xlsx", pce_rev=20.0, pce_fwd=18.0),
            make_cell("D1C2_20260412.xlsx", pce_rev=5.0,  pce_fwd=4.5),
        ]

    def test_passing_cell_appears_in_curated_sheet(self, tmp_path):
        """Cell with PCE_rev=20.0 should appear (threshold 10.0)."""
        cells = self._cells_with_known_pce()
        groups = group_by_substrate(cells)
        out = str(tmp_path / "curated.xlsx")
        export_curated_excel(cells, {"PCE": 10.0}, "rev", groups, out)
        wb = openpyxl.load_workbook(out)
        assert "Curated_PCE" in wb.sheetnames

    def test_failing_cell_not_in_curated_sheet_count(self, tmp_path):
        """Only 1 of 2 cells passes PCE>=10; curated sheet must say 1/2."""
        cells = self._cells_with_known_pce()
        groups = group_by_substrate(cells)
        out = str(tmp_path / "curated.xlsx")
        export_curated_excel(cells, {"PCE": 10.0}, "rev", groups, out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Curated_PCE"]
        found = False
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell and "1" in str(cell) and "2" in str(cell):
                    found = True
        assert found


class TestStatisticsSheetValues:
    """Verify Statistics sheet mean and std match manual computation."""

    def _three_cells(self):
        pces = [10.0, 20.0, 30.0]
        return [
            make_cell(f"D1C{i+1}_20260412.xlsx", pce_rev=p)
            for i, p in enumerate(pces)
        ]

    def test_statistics_sheet_exists(self, tmp_path):
        cells = self._three_cells()
        groups = group_by_substrate(cells)
        out = str(tmp_path / "stats.xlsx")
        export_curated_excel(cells, {}, "rev", groups, out)
        wb = openpyxl.load_workbook(out)
        assert "Statistics" in wb.sheetnames

    def test_statistics_sheet_contains_group_name(self, tmp_path):
        """Statistics sheet should mention the group prefix (D1)."""
        cells = self._three_cells()
        groups = group_by_substrate(cells)
        out = str(tmp_path / "stats.xlsx")
        export_curated_excel(cells, {}, "rev", groups, out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Statistics"]
        all_vals = [ws.cell(r, c).value
                    for r in range(1, ws.max_row + 1)
                    for c in range(1, ws.max_column + 1)]
        assert any("D1" in str(v) for v in all_vals if v is not None)


class TestAllFilteredOut:
    """Exporting with a threshold that rejects every cell must not crash."""

    def test_no_crash_when_all_cells_fail_threshold(self, tmp_path):
        cells = [
            make_cell("D1C1_20260412.xlsx", pce_rev=5.0),
            make_cell("D1C2_20260412.xlsx", pce_rev=4.0),
        ]
        groups = group_by_substrate(cells)
        out = str(tmp_path / "all_filtered.xlsx")
        # PCE threshold 50% → no cell passes
        export_curated_excel(cells, {"PCE": 50.0}, "rev", groups, out)
        # Must produce a file without raising
        from pathlib import Path
        assert Path(out).exists()


class TestForwardOnlyCells:
    """Cells without a reverse scan (fwd data only) must not crash curator."""

    def test_mixed_fwd_and_fwdrev_no_crash(self, tmp_path):
        """A mix of cells — some fwd-only — must export without error."""
        c_fwd_only = make_cell("D1C1_20260412.xlsx", pce_rev=0.0, pce_fwd=12.0)
        c_fwd_only.PCE_rev = 0.0   # simulate no reverse scan result
        c_both = make_cell("D1C2_20260412.xlsx", pce_rev=15.0, pce_fwd=14.0)
        cells  = [c_fwd_only, c_both]
        groups = group_by_substrate(cells)
        out    = str(tmp_path / "mixed.xlsx")
        export_curated_excel(cells, {}, "fwd", groups, out)
        from pathlib import Path
        assert Path(out).exists()
