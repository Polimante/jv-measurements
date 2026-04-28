"""
Tests for app/io/exporter.py

Strategy: build synthetic measurement data (no instrument), call save_data()
and save_dark_jv_data(), then read the resulting xlsx back with openpyxl and
verify that values round-trip correctly.
"""

import math
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from app.io.exporter import save_data, save_dark_jv_data


# ── helpers ───────────────────────────────────────────────────────────────────

IDN = "KEITHLEY INSTRUMENTS INC.,MODEL 2400,1234567,C32"


def _arr(start, stop, n):
    return np.linspace(start, stop, n)


def call_save_data(export_folder, name="TestCell", n_pts=50,
                   PCE_r=15.5, PCE_f=13.9,
                   Voc_r=1.05, Voc_f=1.02,
                   Jsc_r=20.5, Jsc_f=19.8,
                   FF_r=72.1,  FF_f=68.5,
                   Rs_r=3.8,   Rs_f=4.2,
                   Rsh_r=480., Rsh_f=450.,
                   HI=11.0, cell_area=0.07,
                   delay_ms=100, nplc=1.0,
                   scan_rate=250.0):
    """Call save_data() with the correct positional signature."""
    v  = _arr(-0.1, 1.1, n_pts)
    jf = np.array([20.0 - 20.0 * max(0.0, vi / 1.05) ** 2 for vi in v])
    jr = np.array([21.0 - 21.0 * max(0.0, vi / 1.08) ** 2 for vi in v])
    return save_data(
        export_folder, name,
        v, jf, v, jr, 100.0,
        Voc_f, Jsc_f, FF_f, PCE_f, Rs_f, Rsh_f,
        Voc_r, Jsc_r, FF_r, PCE_r, Rs_r, Rsh_r, HI,
        cell_area, delay_ms, nplc, False,
        True, 1.2, 2.0,
        IDN, scan_rate,
    )


def call_save_dark(export_folder, name="TestCell_Dark", n_pts=40):
    """Call save_dark_jv_data() with the correct positional signature."""
    v = _arr(-0.1, 1.0, n_pts)
    j = np.array([-1e-5 * (math.exp(max(vi, -10) / 0.05) - 1) for vi in v])
    return save_dark_jv_data(
        export_folder, name,
        v, j, v, j,
        0.07, 200, 5.0, False,
        IDN,
    )


# ─────────────────────────────────────────────────────────────────────────────
# save_data — illuminated J-V
# ─────────────────────────────────────────────────────────────────────────────

class TestSaveData:

    def test_file_created(self, tmp_path):
        path = call_save_data(str(tmp_path))
        assert Path(path).exists()

    def test_file_is_xlsx(self, tmp_path):
        path = call_save_data(str(tmp_path))
        assert path.endswith(".xlsx")

    def test_sheet_name(self, tmp_path):
        path = call_save_data(str(tmp_path))
        wb = openpyxl.load_workbook(path)
        assert "MeasurementData" in wb.sheetnames

    def test_row1_has_headers(self, tmp_path):
        path = call_save_data(str(tmp_path))
        wb = openpyxl.load_workbook(path)
        ws = wb["MeasurementData"]
        header_row = [ws.cell(1, c).value for c in range(1, 10)]
        assert any(isinstance(h, str) for h in header_row)

    def test_pce_rev_stored_in_row2(self, tmp_path):
        path = call_save_data(str(tmp_path), PCE_r=15.5)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["MeasurementData"]
        # PCE_rev is in column Q (col 17, 1-indexed)
        val = ws.cell(2, 17).value
        assert val == pytest.approx(15.5, abs=0.01)

    def test_pce_fwd_stored_in_row2(self, tmp_path):
        path = call_save_data(str(tmp_path), PCE_f=13.9)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["MeasurementData"]
        # PCE_fwd is in column K (col 11, 1-indexed)
        val = ws.cell(2, 11).value
        assert val == pytest.approx(13.9, abs=0.01)

    def test_hi_stored_in_row2(self, tmp_path):
        path = call_save_data(str(tmp_path), HI=11.0)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["MeasurementData"]
        # HI is in column T (col 20)
        val = ws.cell(2, 20).value
        assert val == pytest.approx(11.0, abs=0.1)

    def test_voltage_array_in_rows3plus(self, tmp_path):
        n = 10
        path = call_save_data(str(tmp_path), n_pts=n)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["MeasurementData"]
        v_read = [ws.cell(3 + i, 1).value for i in range(n)]
        assert all(v is not None for v in v_read)

    def test_sample_name_in_filename(self, tmp_path):
        path = call_save_data(str(tmp_path), name="MySample")
        assert "MySample" in Path(path).name

    def test_scan_rate_stored(self, tmp_path):
        path = call_save_data(str(tmp_path), scan_rate=125.0)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["MeasurementData"]
        # Scan rate is in column AA (col 27)
        val = ws.cell(2, 27).value
        assert val == pytest.approx(125.0, abs=1.0)

    def test_cell_area_stored(self, tmp_path):
        path = call_save_data(str(tmp_path), cell_area=0.09)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["MeasurementData"]
        # Cell area is in column G (col 7)
        val = ws.cell(2, 7).value
        assert val == pytest.approx(0.09, abs=0.001)


# ─────────────────────────────────────────────────────────────────────────────
# save_dark_jv_data
# ─────────────────────────────────────────────────────────────────────────────

class TestSaveDarkJVData:

    def test_file_created(self, tmp_path):
        path = call_save_dark(str(tmp_path))
        assert Path(path).exists()

    def test_sheet_name_dark(self, tmp_path):
        path = call_save_dark(str(tmp_path))
        wb = openpyxl.load_workbook(path)
        assert "Dark_JV_Data" in wb.sheetnames

    def test_dark_filename_contains_dark(self, tmp_path):
        path = call_save_dark(str(tmp_path))
        assert "Dark" in Path(path).name

    def test_voltage_array_stored(self, tmp_path):
        n = 5
        v_expected = np.linspace(-0.1, 1.0, n)
        path = call_save_dark(str(tmp_path), n_pts=n)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Dark_JV_Data"]
        v_stored = [ws.cell(3 + i, 1).value for i in range(n)]
        assert all(v is not None for v in v_stored)
        assert v_stored[0] == pytest.approx(v_expected[0], abs=0.001)


# ─────────────────────────────────────────────────────────────────────────────
# Round-trip: write then read back scalars via data_curator.load_cell_data
# ─────────────────────────────────────────────────────────────────────────────

class TestRoundTrip:
    """
    Write an illuminated J-V Excel file with save_data(), then read it back
    with load_cell_data() and verify scalar parameters survive the round-trip.
    """

    def test_pce_roundtrip(self, tmp_path):
        from app.io.data_curator import load_cell_data

        path = call_save_data(str(tmp_path), PCE_r=16.2, PCE_f=14.7)
        cell = load_cell_data(path)
        assert cell is not None
        assert cell.PCE_rev == pytest.approx(16.2, abs=0.01)
        assert cell.PCE_fwd == pytest.approx(14.7, abs=0.01)

    def test_voc_roundtrip(self, tmp_path):
        from app.io.data_curator import load_cell_data

        path = call_save_data(str(tmp_path), Voc_r=1.06, Voc_f=1.03)
        cell = load_cell_data(path)
        assert cell.Voc_rev == pytest.approx(1.06, abs=0.001)
        assert cell.Voc_fwd == pytest.approx(1.03, abs=0.001)

    def test_hi_roundtrip(self, tmp_path):
        from app.io.data_curator import load_cell_data

        path = call_save_data(str(tmp_path), HI=8.3)
        cell = load_cell_data(path)
        assert cell.HI == pytest.approx(8.3, abs=0.1)

    def test_all_scalars_roundtrip(self, tmp_path):
        from app.io.data_curator import load_cell_data

        path = call_save_data(
            str(tmp_path),
            PCE_r=15.5, PCE_f=13.9,
            Voc_r=1.05, Voc_f=1.02,
            Jsc_r=20.5, Jsc_f=19.8,
            FF_r=72.1,  FF_f=68.5,
            Rs_r=3.8,   Rs_f=4.2,
            Rsh_r=480.0, Rsh_f=450.0,
            HI=11.0,
        )
        cell = load_cell_data(path)
        expected = {
            "PCE_rev": 15.5, "PCE_fwd": 13.9,
            "Voc_rev": 1.05, "Voc_fwd": 1.02,
            "Jsc_rev": 20.5, "Jsc_fwd": 19.8,
            "FF_rev":  72.1, "FF_fwd":  68.5,
            "Rs_rev":   3.8, "Rs_fwd":   4.2,
            "Rsh_rev": 480.0, "Rsh_fwd": 450.0,
            "HI":      11.0,
        }
        for attr, exp in expected.items():
            val = getattr(cell, attr)
            assert val == pytest.approx(exp, abs=0.1), \
                f"{attr}: expected {exp}, got {val}"

    def test_nonexistent_file_returns_none(self):
        from app.io.data_curator import load_cell_data
        result = load_cell_data("/nonexistent/path/file.xlsx")
        assert result is None

    def test_wrong_sheet_returns_none(self, tmp_path):
        from app.io.data_curator import load_cell_data
        import openpyxl

        wb = openpyxl.Workbook()
        wb.active.title = "WrongSheet"
        path = str(tmp_path / "wrong.xlsx")
        wb.save(path)

        result = load_cell_data(path)
        assert result is None
