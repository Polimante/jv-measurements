"""Tests for append_dark_ideality_fit_to_file() in app/io/exporter.py.

Verifies:
  - Sheet is created in an existing xlsx file
  - All values round-trip faithfully
  - Existing DarkIdealityFit sheet is overwritten cleanly
  - Original data sheet is preserved
"""

from __future__ import annotations

import math
import tempfile
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from app.measurement.analysis import DarkIdealityFit
from app.io.exporter import (
    DARK_IDEALITY_SHEET,
    save_dark_jv_data,
    append_dark_ideality_fit_to_file,
)
from app.gui.data_viewer import _read_dark_ideality_fit


# ── Helpers ───────────────────────────────────────────────────────────────────

def _make_fit(**kwargs) -> DarkIdealityFit:
    defaults = dict(
        n_idd=1.55,
        n_idd_err=0.012,
        J_s=3.21e-13,
        J_s_err=4.5e-15,
        r_squared=0.99981,
        rmse_log=0.00312,
        v_min=0.40,
        v_max=1.00,
        n_points=42,
        temperature_K=298.15,
        method="log_linear_regression",
        timestamp="2026-05-11 10:00:00",
    )
    defaults.update(kwargs)
    return DarkIdealityFit(**defaults)


def _save_dark_file(tmp_path) -> str:
    """Save a minimal dark J-V file and return its path."""
    v = np.linspace(0.0, 1.2, 30)
    j = np.zeros_like(v)
    return save_dark_jv_data(
        str(tmp_path), "DarkTest",
        v, j, v, j,
        0.07, 100, 1.0, False, "TEST_IDN",
    )


# ── Sheet creation ────────────────────────────────────────────────────────────

class TestSheetCreation:
    def test_append_creates_dark_ideality_sheet(self, tmp_path):
        fp  = _save_dark_file(tmp_path)
        fit = _make_fit()
        append_dark_ideality_fit_to_file(fp, fit)
        wb = openpyxl.load_workbook(fp)
        assert DARK_IDEALITY_SHEET in wb.sheetnames
        wb.close()

    def test_original_dark_sheet_still_present(self, tmp_path):
        fp  = _save_dark_file(tmp_path)
        append_dark_ideality_fit_to_file(fp, _make_fit())
        wb = openpyxl.load_workbook(fp)
        assert "Dark_JV_Data" in wb.sheetnames
        wb.close()


# ── Round-trip via _read_dark_ideality_fit ────────────────────────────────────

class TestRoundTrip:
    def _write_and_read(self, tmp_path, fit):
        fp = _save_dark_file(tmp_path)
        append_dark_ideality_fit_to_file(fp, fit)
        wb = openpyxl.load_workbook(fp, data_only=True)
        result = _read_dark_ideality_fit(wb[DARK_IDEALITY_SHEET])
        wb.close()
        return result

    def test_n_idd_roundtrip(self, tmp_path):
        fit = _make_fit(n_idd=1.55)
        r = self._write_and_read(tmp_path, fit)
        assert r["n_idd"] == pytest.approx(1.55, rel=1e-4)

    def test_n_idd_err_roundtrip(self, tmp_path):
        fit = _make_fit(n_idd_err=0.012)
        r = self._write_and_read(tmp_path, fit)
        assert r["n_idd_err"] == pytest.approx(0.012, rel=1e-3)

    def test_Js_roundtrip(self, tmp_path):
        fit = _make_fit(J_s=3.21e-13)
        r = self._write_and_read(tmp_path, fit)
        assert r["J_s"] == pytest.approx(3.21e-13, rel=1e-3)

    def test_Js_err_roundtrip(self, tmp_path):
        fit = _make_fit(J_s_err=4.5e-15)
        r = self._write_and_read(tmp_path, fit)
        assert r["J_s_err"] == pytest.approx(4.5e-15, rel=1e-2)

    def test_r_squared_roundtrip(self, tmp_path):
        fit = _make_fit(r_squared=0.99981)
        r = self._write_and_read(tmp_path, fit)
        assert r["r_squared"] == pytest.approx(0.99981, rel=1e-4)

    def test_temperature_roundtrip(self, tmp_path):
        fit = _make_fit(temperature_K=305.0)
        r = self._write_and_read(tmp_path, fit)
        assert r["temperature_K"] == pytest.approx(305.0, rel=1e-4)

    def test_method_roundtrip(self, tmp_path):
        r = self._write_and_read(tmp_path, _make_fit())
        assert r["method"] == "log_linear_regression"

    def test_n_points_roundtrip(self, tmp_path):
        fit = _make_fit(n_points=37)
        r = self._write_and_read(tmp_path, fit)
        assert r["n_points"] == 37


# ── Overwrite behaviour ───────────────────────────────────────────────────────

class TestOverwrite:
    def test_overwrite_existing_DarkIdealityFit_sheet(self, tmp_path):
        fp = _save_dark_file(tmp_path)
        append_dark_ideality_fit_to_file(fp, _make_fit(n_idd=1.40))
        append_dark_ideality_fit_to_file(fp, _make_fit(n_idd=1.80))
        wb = openpyxl.load_workbook(fp, data_only=True)
        result = _read_dark_ideality_fit(wb[DARK_IDEALITY_SHEET])
        wb.close()
        assert result["n_idd"] == pytest.approx(1.80, rel=1e-4)

    def test_only_one_dark_ideality_sheet_after_overwrite(self, tmp_path):
        fp = _save_dark_file(tmp_path)
        append_dark_ideality_fit_to_file(fp, _make_fit())
        append_dark_ideality_fit_to_file(fp, _make_fit())
        wb = openpyxl.load_workbook(fp)
        count = wb.sheetnames.count(DARK_IDEALITY_SHEET)
        wb.close()
        assert count == 1
