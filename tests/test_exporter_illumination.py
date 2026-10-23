"""The irradiance the PCE was divided by must be in the saved J-V file.

Before 2026-10 ``save_data()`` received the illumination but never wrote it,
so a MeasurementData file could not show — or let anyone recompute — its PCE.
It now lives in fixed trailing columns AB-AE, read back by the Data Viewer.
"""

from __future__ import annotations

import numpy as np
import openpyxl
import pytest

from app.gui.data_viewer import load_jv_file
from app.io.exporter import save_data

AREA = 0.06


def _save(folder, illum=100.0, **kw):
    v = np.linspace(-0.1, 1.0, 12)
    j = 18 - 18 * np.clip(v / 0.8, 0, None) ** 6
    return save_data(
        str(folder), "Cell", v, j, v[::-1], j[::-1], illum,
        0.8, 18, 70, 10.08, 5, 1e4, 0.8, 18, 70, 10.08, 5, 1e4, 0.0,
        AREA, 33.3, 1.0, False, False, 0, 0, "IDN", **kw)


def test_trailing_columns_have_fixed_headers(tmp_path):
    ws = openpyxl.load_workbook(_save(tmp_path))["MeasurementData"]
    assert [ws.cell(1, c).value for c in range(27, 32)] == [
        "Scan Rate (mV/s)", "Scan Order", "Illumination (mW/cm²)",
        "Detector Power (mW)", "Detector Area (cm²)"]


def test_illumination_is_always_written(tmp_path):
    ws = openpyxl.load_workbook(_save(tmp_path, illum=87.5))["MeasurementData"]
    assert ws.cell(2, 29).value == pytest.approx(87.5)
    # Optional values left out keep their cells empty, positions unchanged.
    assert ws.cell(2, 28).value is None
    assert ws.cell(2, 30).value is None


def test_round_trip_through_the_data_viewer(tmp_path):
    fp = _save(tmp_path, illum=100.0, scan_order="REV->FWD",
               detector_power=113.0, detector_area=1.13)
    meta = load_jv_file(fp)
    assert meta["illumination"] == pytest.approx(100.0)
    assert meta["detector_power"] == pytest.approx(113.0)
    assert meta["detector_area"] == pytest.approx(1.13)
    assert meta["scan_order"] == "REV->FWD"
    # The columns before the new ones did not move.
    assert meta["scan_rate"] is not None
    assert meta["PCE_r"] == pytest.approx(10.08)


def test_pce_can_be_recomputed_from_the_file(tmp_path):
    meta = load_jv_file(_save(tmp_path, illum=100.0))
    pce = meta["Voc_r"] * meta["Jsc_r"] * meta["FF_r"] / meta["illumination"]
    assert pce == pytest.approx(meta["PCE_r"])


def test_old_file_without_the_columns_still_opens(tmp_path):
    fp = _save(tmp_path)
    wb = openpyxl.load_workbook(fp)
    wb["MeasurementData"].delete_cols(28, 4)     # what a pre-2026-10 file looks like
    wb.save(fp)
    meta = load_jv_file(fp)
    assert meta["illumination"] is None
    assert meta["scan_order"] is None
    assert meta["cell_area"] == pytest.approx(AREA)


# ── Data Viewer panel ────────────────────────────────────────────────────────

def _panel_text(widget):
    """Every Label text under *widget*, in packing order."""
    out = []
    for child in widget.winfo_children():
        try:
            if child.winfo_class() in ("Label", "Labelframe"):
                out.append(str(child.cget("text")))
        except Exception:
            pass
        out.extend(_panel_text(child))
    return out


@pytest.mark.gui
def test_viewer_shows_the_illumination_section(tk_root, tmp_path):
    from app.gui.data_viewer import DataViewerWindow

    fp = _save(tmp_path, illum=100.0, scan_order="REV->FWD",
               detector_power=113.0, detector_area=1.13)
    win = DataViewerWindow(tk_root, fp)
    try:
        text = _panel_text(win._meta_inner)
    finally:
        win.win.destroy()
    assert "Illumination" in text
    assert "100.00 mW/cm²" in text
    assert "1.000 sun" in text
    assert "113.000 mW" in text
    assert "REV->FWD" in text


@pytest.mark.gui
def test_viewer_says_so_when_an_old_file_has_no_illumination(tk_root, tmp_path):
    from app.gui.data_viewer import DataViewerWindow

    fp = _save(tmp_path)
    wb = openpyxl.load_workbook(fp)
    wb["MeasurementData"].delete_cols(28, 4)
    wb.save(fp)
    win = DataViewerWindow(tk_root, fp)
    try:
        text = _panel_text(win._meta_inner)
    finally:
        win.win.destroy()
    assert "not recorded (older file)" in text
