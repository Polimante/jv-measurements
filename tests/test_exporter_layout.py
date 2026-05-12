"""
Exporter layout tests — Part B of the bug-hunt suite.

For each measurement type that produces an xlsx file the test:
  1. Calls the exporter with known scalars and arrays.
  2. Reopens the file with openpyxl.
  3. Asserts every header and scalar is in the expected cell.
  4. Verifies V/J arrays start at the expected row.
  5. Round-trips a scalar value: what was written equals what is read back.

openpyxl row/col indices are 1-based throughout.
xlsxwriter column-0 maps to openpyxl column-1.
"""
from __future__ import annotations

import numpy as np
import openpyxl
import pytest

from app.io.exporter import (
    save_data,
    save_dark_jv_data,
    save_stab_current_data,
    save_pure_mpp_data,
    save_mpp_data,
)


# ── Shared test data ──────────────────────────────────────────────────────────

_AREA   = 0.07
_ILLUM  = 100.0
_N      = 30           # data-array length
_V_F    = np.linspace(-0.05, 1.1,  _N)
_J_F    = 20.0 * (1.0 - _V_F / 1.05)
_V_R    = _V_F.copy()
_J_R    = _J_F * 1.02  # slightly different so fwd ≠ rev


def _save_illuminated(tmp_path):
    """Write a minimal illuminated J-V file; return its path."""
    return save_data(
        export_folder=str(tmp_path),
        sample_name="TestCell",
        v_f=_V_F, j_f=_J_F,
        v_r=_V_R, j_r=_J_R,
        illum=_ILLUM,
        Voc_f=0.950, Jsc_f=19.50, FF_f=65.0, PCE_f=12.05,
        Rs_f=3.5,   Rsh_f=450.0,
        Voc_r=1.000, Jsc_r=20.00, FF_r=70.0, PCE_r=14.00,
        Rs_r=3.0,   Rsh_r=500.0,
        HI=13.9,
        cell_area=_AREA, delay_ms=100.0, nplc=1.0, four_wire=False,
        prebias_enabled=False, prebias_voltage=0.0, prebias_duration=0.0,
        instrument_idn="Keithley2400",
        scan_rate=None,
        fit_result=None,
    )


# ── Illuminated J-V ───────────────────────────────────────────────────────────

class TestIlluminatedJVLayout:

    def setup_method(self, _method):
        pass

    def test_sheet_name(self, tmp_path):
        fp = _save_illuminated(tmp_path)
        wb = openpyxl.load_workbook(fp)
        assert "MeasurementData" in wb.sheetnames

    def test_header_row_not_empty(self, tmp_path):
        fp = _save_illuminated(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["MeasurementData"]
        # Row 1 must be headers; at least the first header must be non-empty
        assert ws.cell(1, 1).value is not None
        assert "Voltage" in str(ws.cell(1, 1).value)

    def test_scalar_voc_rev_position(self, tmp_path):
        """Voc_rev is written to row 2, col 14 (N) per exporter.py."""
        fp = _save_illuminated(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["MeasurementData"]
        val = ws.cell(2, 14).value
        assert val == pytest.approx(1.000, abs=1e-4)

    def test_scalar_pce_rev_position(self, tmp_path):
        """PCE_rev is at row 2, col 17 (Q)."""
        fp = _save_illuminated(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["MeasurementData"]
        val = ws.cell(2, 17).value
        assert val == pytest.approx(14.00, abs=1e-4)

    def test_scalar_cell_area_position(self, tmp_path):
        """Cell area is at row 2, col 7 (G)."""
        fp = _save_illuminated(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["MeasurementData"]
        val = ws.cell(2, 7).value
        assert val == pytest.approx(_AREA, abs=1e-6)

    def test_scalar_hi_position(self, tmp_path):
        """HI is at row 2, col 20 (T)."""
        fp = _save_illuminated(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["MeasurementData"]
        val = ws.cell(2, 20).value
        assert val == pytest.approx(13.9, abs=1e-4)

    def test_vj_arrays_start_at_row_3(self, tmp_path):
        """Forward voltage array must start at row 3, col 1 (A)."""
        fp = _save_illuminated(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["MeasurementData"]
        first_v = ws.cell(3, 1).value
        assert first_v is not None
        assert abs(float(first_v) - float(_V_F[0])) < 1e-5

    def test_reverse_vj_arrays_start_at_row_3(self, tmp_path):
        """Reverse voltage array must start at row 3, col 4 (D)."""
        fp = _save_illuminated(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["MeasurementData"]
        first_v_r = ws.cell(3, 4).value
        assert first_v_r is not None
        assert abs(float(first_v_r) - float(_V_R[0])) < 1e-5

    def test_round_trip_pce_fwd(self, tmp_path):
        """Written PCE_fwd (col 11) reads back within float precision."""
        fp = _save_illuminated(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["MeasurementData"]
        val = ws.cell(2, 11).value
        assert val == pytest.approx(12.05, rel=1e-5)

    def test_header_style_bold(self, tmp_path):
        """Header cells must be bold (style check)."""
        fp = _save_illuminated(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["MeasurementData"]
        assert ws.cell(1, 1).font.bold is True

    def test_header_background_blue(self, tmp_path):
        """Header background must contain the project blue #1565C0."""
        fp = _save_illuminated(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["MeasurementData"]
        rgb = ws.cell(1, 1).fill.fgColor.rgb
        assert "1565C0" in rgb.upper()


# ── Dark J-V ─────────────────────────────────────────────────────────────────

class TestDarkJVLayout:

    def _write(self, tmp_path):
        return save_dark_jv_data(
            export_folder=str(tmp_path),
            sample_name="TestDark",
            v_f=_V_F, j_f=-np.exp(_V_F / 0.05) * 1e-3,
            v_r=_V_R, j_r=-np.exp(_V_R / 0.05) * 1e-3,
            cell_area=_AREA, delay_ms=100.0, nplc=1.0,
            four_wire=False, instrument_idn="Keithley2400",
        )

    def test_dark_sheet_name(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        assert "Dark_JV_Data" in wb.sheetnames

    def test_dark_header_row(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["Dark_JV_Data"]
        assert ws.cell(1, 1).value is not None
        assert "Voltage" in str(ws.cell(1, 1).value)

    def test_dark_cell_area_position(self, tmp_path):
        """Cell area is at row 2, col 5 (E) for dark J-V."""
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["Dark_JV_Data"]
        val = ws.cell(2, 5).value
        assert val == pytest.approx(_AREA, abs=1e-6)

    def test_dark_arrays_start_at_row_3(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["Dark_JV_Data"]
        assert ws.cell(3, 1).value is not None


# ── Stabilized Current ────────────────────────────────────────────────────────

class TestStabCurrentLayout:

    def _write(self, tmp_path):
        times = np.linspace(0, 120, 25)
        j_vals  = np.full(25, 18.5)
        pce_vals = np.full(25, 14.2)
        return save_stab_current_data(
            export_folder=str(tmp_path),
            sample_name="TestStab",
            times=times, j_vals=j_vals, pce_vals=pce_vals,
            cell_area=_AREA, v_applied=0.82, illumination=_ILLUM,
            duration_s=120.0, interval_ms=5000,
            threshold_pct_per_min=0.01,
            stabilized_at=60.0,
        )

    def test_stab_sheet_name(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        assert "StabCurrent_Data" in wb.sheetnames

    def test_stab_time_header(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["StabCurrent_Data"]
        assert "Time" in str(ws.cell(1, 1).value)

    def test_stab_arrays_start_row_3(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["StabCurrent_Data"]
        assert ws.cell(3, 1).value is not None


# ── Pure MPP Tracking ─────────────────────────────────────────────────────────

class TestPureMPPLayout:

    def _write(self, tmp_path):
        t = np.linspace(0, 600, 50)
        v = np.full(50, 0.85)
        j = np.full(50, 18.0)
        pce = np.full(50, 15.3)
        return save_pure_mpp_data(
            export_folder=str(tmp_path),
            sample_name="TestPureMPP",
            time_s=t, v_track=v, j_track=j, pce_track=pce,
            cell_area=_AREA, illumination=_ILLUM,
            nplc=1.0, delay_ms=100.0,
        )

    def test_pure_mpp_sheet_name(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        assert "PureMPP_Data" in wb.sheetnames

    def test_pure_mpp_time_header(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["PureMPP_Data"]
        assert "Time" in str(ws.cell(1, 1).value)

    def test_pure_mpp_arrays_start_row_3(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["PureMPP_Data"]
        assert ws.cell(3, 1).value is not None


# ── MPP Tracking (JV-SPO) ────────────────────────────────────────────────────

class TestMPPLayout:

    def _cycle(self, n):
        return {
            "cycle": n,
            "time_s": n * 600.0,
            "Vmpp": 0.84, "Jmpp": 18.5,
            "Voc": 1.05,  "Jsc": 20.0,
            "FF": 70.0,   "PCE": 14.7,
            "HI": 8.0,
            "v_rev": _V_R, "j_rev": _J_R,
            "v_fwd": None, "j_fwd": None,
        }

    def _write(self, tmp_path):
        cycles = [self._cycle(1), self._cycle(2)]
        t = np.linspace(0, 1200, 60)
        j = np.full(60, 18.0)
        pce = np.full(60, 14.5)
        idx = [1] * 30 + [2] * 30
        return save_mpp_data(
            export_folder=str(tmp_path),
            sample_name="TestMPP",
            cycles=cycles,
            spo_time=list(t), spo_j=list(j), spo_pce=list(pce),
            spo_cycle_idx=idx,
            cell_area=_AREA, illumination=_ILLUM,
            nplc=1.0, delay_ms=100.0,
            spo_duration_s=60.0, total_duration_min=20.0,
        )

    def test_mpp_spo_sheet_exists(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        assert "MPP_SPO" in wb.sheetnames

    def test_mpp_jv_sheet_exists(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        assert "MPP_JV" in wb.sheetnames

    def test_mpp_cycle_sheet_exists(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        assert "JV_Cycle_1" in wb.sheetnames

    def test_mpp_jv_header_row(self, tmp_path):
        fp = self._write(tmp_path)
        wb = openpyxl.load_workbook(fp)
        ws = wb["MPP_JV"]
        assert ws.cell(1, 1).value is not None
        assert "Cycle" in str(ws.cell(1, 1).value)
