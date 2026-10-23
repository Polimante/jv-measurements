"""Tests for SingleDiodeFit sheet in app/io/exporter.py.

Verifies that:
  - save_data() creates a SingleDiodeFit sheet when fit_result is supplied.
  - save_data() creates no such sheet when fit_result is None.
  - Parameters survive a write-then-read round-trip via data_viewer's reader.
  - append_single_diode_fit_to_file() correctly adds/replaces the sheet.
"""

from __future__ import annotations

import math
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from app.measurement.analysis import SingleDiodeFit
from app.io.exporter import FIT_SHEET, save_data, append_single_diode_fit_to_file
from app.gui.data_viewer import _read_single_diode_fit


# ── Helpers ───────────────────────────────────────────────────────────────────

IDN = "KEITHLEY INSTRUMENTS INC.,MODEL 2400,1234567,C32"

_KNOWN_PARAMS = {
    "J_ph": 20.5,
    "J0":   1.23e-11,
    "n":    1.75,
    "Rs":   4.2,
    "Rsh":  3000.0,
}
_KNOWN_ERRORS = {
    "J_ph": 0.01,
    "J0":   2.1e-13,
    "n":    0.03,
    "Rs":   0.2,
    "Rsh":  100.0,
}

def _make_fit(**kwargs) -> SingleDiodeFit:
    defaults = dict(
        parameters=_KNOWN_PARAMS.copy(),
        std_errors=_KNOWN_ERRORS.copy(),
        r_squared=0.99982,
        rmse=0.0123,
        rmse_log=0.0045,
        converged=True,
        method="curve_fit+brentq",
        bounded_parameters=[],
        temperature_K=298.15,
        timestamp="2026-05-05 12:00:00",
    )
    defaults.update(kwargs)
    return SingleDiodeFit(**defaults)


def _save(tmp_path, fit=None, n_pts=20):
    """Call save_data with minimal synthetic data and an optional fit."""
    v = np.linspace(-0.1, 1.1, n_pts)
    j = np.array([20.0 - 20.0 * max(0.0, vi / 1.05) ** 2 for vi in v])
    return save_data(
        str(tmp_path), "TestSample",
        v, j, v, j, 100.0,
        1.02, 19.5, 68.0, 13.3, 4.0, 450.0,
        1.05, 20.0, 72.0, 15.0, 3.5, 480.0,
        10.0,
        0.07, 100, 1.0, False,
        True, 1.2, 2.0,
        IDN,
        fit_result=fit,
    )


# ── Sheet creation ────────────────────────────────────────────────────────────

class TestSheetCreation:
    def test_fit_sheet_present_when_fit_supplied(self, tmp_path):
        fp = _save(tmp_path, fit=_make_fit())
        wb = openpyxl.load_workbook(fp)
        assert FIT_SHEET in wb.sheetnames
        wb.close()

    def test_fit_sheet_absent_when_fit_none(self, tmp_path):
        fp = _save(tmp_path, fit=None)
        wb = openpyxl.load_workbook(fp)
        assert FIT_SHEET not in wb.sheetnames
        wb.close()

    def test_measurement_sheet_still_present(self, tmp_path):
        fp = _save(tmp_path, fit=_make_fit())
        wb = openpyxl.load_workbook(fp)
        assert "MeasurementData" in wb.sheetnames
        wb.close()


# ── Round-trip: write via save_data, read via _read_single_diode_fit ─────────

class TestRoundTrip:
    def test_parameters_roundtrip(self, tmp_path):
        fit = _make_fit()
        fp  = _save(tmp_path, fit=fit)
        wb  = openpyxl.load_workbook(fp, data_only=True)
        result = _read_single_diode_fit(wb[FIT_SHEET])
        wb.close()

        for name, expected in _KNOWN_PARAMS.items():
            got = result["parameters"].get(name)
            assert got is not None, f"Missing parameter {name}"
            assert got == pytest.approx(expected, rel=1e-4), \
                f"{name}: expected {expected}, got {got}"

    def test_std_errors_roundtrip(self, tmp_path):
        fit = _make_fit()
        fp  = _save(tmp_path, fit=fit)
        wb  = openpyxl.load_workbook(fp, data_only=True)
        result = _read_single_diode_fit(wb[FIT_SHEET])
        wb.close()

        for name, expected in _KNOWN_ERRORS.items():
            got = result["std_errors"].get(name)
            assert got is not None, f"Missing std_error for {name}"
            assert got == pytest.approx(expected, rel=1e-3), \
                f"Std error {name}: expected {expected}, got {got}"

    def test_r_squared_roundtrip(self, tmp_path):
        fit = _make_fit(r_squared=0.99982)
        fp  = _save(tmp_path, fit=fit)
        wb  = openpyxl.load_workbook(fp, data_only=True)
        result = _read_single_diode_fit(wb[FIT_SHEET])
        wb.close()
        assert result["r_squared"] == pytest.approx(0.99982, rel=1e-4)

    def test_rmse_roundtrip(self, tmp_path):
        fit = _make_fit(rmse=0.0123)
        fp  = _save(tmp_path, fit=fit)
        wb  = openpyxl.load_workbook(fp, data_only=True)
        result = _read_single_diode_fit(wb[FIT_SHEET])
        wb.close()
        assert result["rmse"] == pytest.approx(0.0123, rel=1e-3)

    def test_temperature_roundtrip(self, tmp_path):
        fit = _make_fit(temperature_K=305.0)
        fp  = _save(tmp_path, fit=fit)
        wb  = openpyxl.load_workbook(fp, data_only=True)
        result = _read_single_diode_fit(wb[FIT_SHEET])
        wb.close()
        assert result["temperature_K"] == pytest.approx(305.0, rel=1e-4)

    def test_bounded_parameters_roundtrip_nonempty(self, tmp_path):
        fit = _make_fit(bounded_parameters=["n", "Rsh"])
        fp  = _save(tmp_path, fit=fit)
        wb  = openpyxl.load_workbook(fp, data_only=True)
        result = _read_single_diode_fit(wb[FIT_SHEET])
        wb.close()
        assert "n"   in result["bounded_parameters"]
        assert "Rsh" in result["bounded_parameters"]

    def test_bounded_parameters_roundtrip_empty(self, tmp_path):
        fit = _make_fit(bounded_parameters=[])
        fp  = _save(tmp_path, fit=fit)
        wb  = openpyxl.load_workbook(fp, data_only=True)
        result = _read_single_diode_fit(wb[FIT_SHEET])
        wb.close()
        assert result["bounded_parameters"] == []

    def test_method_roundtrip(self, tmp_path):
        fit = _make_fit(method="curve_fit+brentq")
        fp  = _save(tmp_path, fit=fit)
        wb  = openpyxl.load_workbook(fp, data_only=True)
        result = _read_single_diode_fit(wb[FIT_SHEET])
        wb.close()
        assert "brentq" in result["method"]


# ── load_jv_file integration ──────────────────────────────────────────────────

class TestLoadJvFileIntegration:
    def test_load_jv_file_includes_fit(self, tmp_path):
        from app.gui.data_viewer import load_jv_file
        fp  = _save(tmp_path, fit=_make_fit())
        meta = load_jv_file(fp)
        assert "single_diode_fit" in meta
        assert meta["single_diode_fit"] is not None

    def test_load_jv_file_no_fit_key_absent_or_none(self, tmp_path):
        from app.gui.data_viewer import load_jv_file
        fp   = _save(tmp_path, fit=None)
        meta = load_jv_file(fp)
        # Either the key is absent or None — backward-compatible behaviour
        assert meta.get("single_diode_fit") is None


# ── append_single_diode_fit_to_file ──────────────────────────────────────────

class TestAppendToFile:
    def test_append_creates_sheet(self, tmp_path):
        fp  = _save(tmp_path, fit=None)
        fit = _make_fit()
        append_single_diode_fit_to_file(fp, fit)
        wb = openpyxl.load_workbook(fp)
        assert FIT_SHEET in wb.sheetnames
        wb.close()

    def test_append_overwrites_existing_sheet(self, tmp_path):
        fp = _save(tmp_path, fit=_make_fit(r_squared=0.990))
        # Append a second time with updated R²
        append_single_diode_fit_to_file(fp, _make_fit(r_squared=0.999))
        wb = openpyxl.load_workbook(fp, data_only=True)
        result = _read_single_diode_fit(wb[FIT_SHEET])
        wb.close()
        # Should reflect the second write
        assert result["r_squared"] == pytest.approx(0.999, rel=1e-3)

    def test_append_preserves_measurement_sheet(self, tmp_path):
        fp = _save(tmp_path, fit=None)
        append_single_diode_fit_to_file(fp, _make_fit())
        wb = openpyxl.load_workbook(fp)
        assert "MeasurementData" in wb.sheetnames
        wb.close()
