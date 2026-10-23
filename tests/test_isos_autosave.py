"""
Tests for app/io/isos_autosave.py — CSVStream and consolidate_csv_into_xlsx.
No instrument or GUI required.
"""

import csv
import os

import openpyxl
import pytest

from app.io.isos_autosave import CSVStream, consolidate_csv_into_xlsx


_COLUMNS = ["t", "v", "j", "pce"]


def _rows(n):
    return [
        {"t": i * 0.5, "v": 0.8 + i * 1e-3, "j": 19.0 - i * 1e-2, "pce": 14.0 + i * 1e-3}
        for i in range(n)
    ]


# ── CSVStream ───────────────────────────────────────────────────────────────

class TestCSVStream:

    def test_open_writes_header(self, tmp_path):
        path = str(tmp_path / "stream.csv")
        stream = CSVStream(path, _COLUMNS)
        stream.open()
        stream.close()
        with open(path, newline="", encoding="utf-8") as fh:
            header = next(csv.reader(fh))
        assert header == _COLUMNS

    def test_append_100_rows(self, tmp_path):
        path = str(tmp_path / "stream.csv")
        stream = CSVStream(path, _COLUMNS)
        stream.open()
        for row in _rows(100):
            stream.append(row)
        stream.close()
        with open(path, newline="", encoding="utf-8") as fh:
            data = list(csv.reader(fh))
        assert len(data) == 101  # header + 100 rows
        assert data[1][0] == "0.0"
        assert float(data[100][0]) == pytest.approx(99 * 0.5, abs=1e-9)

    def test_append_before_open_raises(self, tmp_path):
        stream = CSVStream(str(tmp_path / "s.csv"), _COLUMNS)
        with pytest.raises(RuntimeError):
            stream.append({"t": 0})

    def test_missing_columns_filled_blank(self, tmp_path):
        path = str(tmp_path / "stream.csv")
        stream = CSVStream(path, _COLUMNS)
        stream.open()
        stream.append({"t": 1.0})  # v, j, pce missing
        stream.close()
        with open(path, newline="", encoding="utf-8") as fh:
            data = list(csv.reader(fh))
        assert data[1] == ["1.0", "", "", ""]

    def test_extra_columns_ignored(self, tmp_path):
        path = str(tmp_path / "stream.csv")
        stream = CSVStream(path, _COLUMNS)
        stream.open()
        stream.append({"t": 1.0, "v": 0.8, "j": 19.0, "pce": 14.0, "extra": "x"})
        stream.close()
        with open(path, newline="", encoding="utf-8") as fh:
            data = list(csv.reader(fh))
        assert len(data[1]) == len(_COLUMNS)

    def test_context_manager(self, tmp_path):
        path = str(tmp_path / "stream.csv")
        with CSVStream(path, _COLUMNS) as stream:
            stream.append({"t": 0.0, "v": 0.8, "j": 19.0, "pce": 14.0})
        with open(path, newline="", encoding="utf-8") as fh:
            data = list(csv.reader(fh))
        assert len(data) == 2


# ── consolidate_csv_into_xlsx ────────────────────────────────────────────────

class TestConsolidate:

    def _make_csv(self, tmp_path, n=100):
        path = str(tmp_path / "stress.csv")
        stream = CSVStream(path, _COLUMNS)
        stream.open()
        rows = _rows(n)
        for row in rows:
            stream.append(row)
        stream.close()
        return path, rows

    def test_creates_sheet_in_fresh_xlsx(self, tmp_path):
        csv_path, _ = self._make_csv(tmp_path)
        xlsx_path = str(tmp_path / "out.xlsx")
        consolidate_csv_into_xlsx(csv_path, xlsx_path, "Stress")
        wb = openpyxl.load_workbook(xlsx_path)
        assert "Stress" in wb.sheetnames
        wb.close()

    def test_all_rows_roundtrip(self, tmp_path):
        csv_path, rows = self._make_csv(tmp_path, n=100)
        xlsx_path = str(tmp_path / "out.xlsx")
        consolidate_csv_into_xlsx(csv_path, xlsx_path, "Stress")
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb["Stress"]
        assert ws.max_row == 101  # header + 100
        header = [ws.cell(1, c).value for c in range(1, len(_COLUMNS) + 1)]
        assert header == _COLUMNS
        # Verify a few values (stored as strings from CSV)
        assert float(ws.cell(2, 1).value) == pytest.approx(rows[0]["t"], abs=1e-9)
        assert float(ws.cell(101, 1).value) == pytest.approx(rows[99]["t"], abs=1e-9)
        wb.close()

    def test_csv_renamed_to_consumed(self, tmp_path):
        csv_path, _ = self._make_csv(tmp_path)
        xlsx_path = str(tmp_path / "out.xlsx")
        consolidate_csv_into_xlsx(csv_path, xlsx_path, "Stress")
        assert not os.path.exists(csv_path)
        assert os.path.exists(csv_path + ".consumed")

    def test_appends_to_existing_xlsx(self, tmp_path):
        # Pre-create xlsx with an existing sheet.
        xlsx_path = str(tmp_path / "out.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Existing"
        wb.active["A1"] = "keep"
        wb.save(xlsx_path)
        wb.close()

        csv_path, _ = self._make_csv(tmp_path)
        consolidate_csv_into_xlsx(csv_path, xlsx_path, "Stress")
        wb = openpyxl.load_workbook(xlsx_path)
        assert "Existing" in wb.sheetnames
        assert "Stress" in wb.sheetnames
        assert wb["Existing"]["A1"].value == "keep"
        wb.close()
