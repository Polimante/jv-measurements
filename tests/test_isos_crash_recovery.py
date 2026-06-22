"""Crash-recovery tests for interrupted ISOS stress runs.

Covers the Tk-free recovery layer in ``app.io.isos_recovery``: recent-folder
persistence, orphan-CSV detection/pairing, consolidation, and discard. Also
simulates a real crash (a runner that dies after streaming but before
consolidating) and verifies the leftover CSV is recoverable, and that the GUI
exposes a recovery-dialog method.
"""

from __future__ import annotations

import csv
import os
import threading
from unittest.mock import patch

import openpyxl
import pytest

from app.io import isos_recovery
from app.io.isos_recovery import (
    add_recent_folder, load_recent_folders,
    find_orphan_stress_csvs, consolidate_orphan, discard_orphan,
)
from app.io.isos_tracking import create_isos_tracking_file
from app.measurement.isos_continuous_runner import (
    ContinuousISOSRunner, _STRESS_CSV_COLUMNS,
)

from tests._isos_runner_helpers import FakeKeithley, VirtualClock, make_runner_config


# ── Recent-folder persistence ───────────────────────────────────────────────

def _patch_recent(tmp_path):
    target = str(tmp_path / "recent_folders.json")
    return patch.object(isos_recovery, "recent_folders_path", lambda: target)


def test_recent_folders_roundtrip(tmp_path):
    with _patch_recent(tmp_path):
        assert load_recent_folders() == []
        a = str(tmp_path / "A")
        b = str(tmp_path / "B")
        add_recent_folder(a)
        add_recent_folder(b)
        folders = load_recent_folders()
        # Most-recent first, both present, absolute, deduped.
        assert os.path.normcase(folders[0]) == os.path.normcase(os.path.abspath(b))
        add_recent_folder(a)                      # re-adding moves to front
        folders = load_recent_folders()
        assert os.path.normcase(folders[0]) == os.path.normcase(os.path.abspath(a))
        assert len(folders) == 2                  # deduped


def test_load_recent_folders_tolerates_corrupt_file(tmp_path):
    with _patch_recent(tmp_path):
        with open(isos_recovery.recent_folders_path(), "w", encoding="utf-8") as fh:
            fh.write("not json {{{")
        assert load_recent_folders() == []


# ── Orphan helpers (hand-built scenario) ────────────────────────────────────

def _make_orphan(folder, sample="CrashCell", n_rows=3):
    """Create a tracking workbook + a paired in-progress stress CSV."""
    xlsx = create_isos_tracking_file(folder, {
        "protocol_id": "ISOS-L-1", "sample_name": sample})
    csv_path = os.path.join(folder, f"{sample}_stress_inprogress_20260512_142033.csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=_STRESS_CSV_COLUMNS)
        w.writeheader()
        for i in range(n_rows):
            w.writerow({"elapsed_h": i * 0.01, "t_s": i * 36.0,
                        "v": 0.8, "j": 19.0 - i, "pce": 15.0 - i,
                        "phase": "stress_oc"})
    return csv_path, xlsx


def test_find_orphan_pairs_csv_with_xlsx(tmp_path):
    folder = str(tmp_path)
    csv_path, xlsx = _make_orphan(folder)
    orphans = find_orphan_stress_csvs([folder])
    assert len(orphans) == 1
    found_csv, found_xlsx = orphans[0]
    assert os.path.samefile(found_csv, csv_path)
    assert os.path.samefile(found_xlsx, xlsx)


def test_unpaired_orphan_excluded_unless_requested(tmp_path):
    folder = str(tmp_path)
    # CSV with no matching xlsx.
    csv_path = os.path.join(folder, "Lonely_stress_inprogress_20260101_000000.csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        csv.DictWriter(fh, fieldnames=_STRESS_CSV_COLUMNS).writeheader()
    assert find_orphan_stress_csvs([folder]) == []
    inc = find_orphan_stress_csvs([folder], include_unpaired=True)
    assert len(inc) == 1 and inc[0][1] is None


def test_consumed_and_discarded_files_are_not_orphans(tmp_path):
    folder = str(tmp_path)
    csv_path, _ = _make_orphan(folder)
    os.replace(csv_path, csv_path + ".consumed")
    assert find_orphan_stress_csvs([folder]) == []


def test_consolidate_orphan_adds_sheet_and_consumes(tmp_path):
    folder = str(tmp_path)
    csv_path, xlsx = _make_orphan(folder, n_rows=4)
    sheet = consolidate_orphan(csv_path, xlsx)

    wb = openpyxl.load_workbook(xlsx)
    assert sheet in wb.sheetnames
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    assert len(rows) == 5                          # header + 4 data rows
    assert not os.path.exists(csv_path)
    assert os.path.exists(csv_path + ".consumed")
    # After consolidation it is no longer detected as an orphan.
    assert find_orphan_stress_csvs([folder]) == []


def test_discard_orphan_preserves_file(tmp_path):
    folder = str(tmp_path)
    csv_path, _ = _make_orphan(folder)
    new = discard_orphan(csv_path)
    assert new == csv_path + ".discarded"
    assert os.path.exists(new)
    assert not os.path.exists(csv_path)
    assert find_orphan_stress_csvs([folder]) == []


# ── True crash simulation: runner dies before consolidation ─────────────────

def test_crashed_runner_leaves_recoverable_csv(tmp_path):
    cfg = make_runner_config(
        tmp_path,
        clock=VirtualClock(),
        total_duration_h=0.02,        # ~2 scans worth of stress
        scan_interval_min=0.6,
        stress_log_interval_s=6.0,
        baseline_n_scans=0,
        spo_duration_s=0.0,
        pre_scan_relaxation_s=0.0,
    )
    fake = FakeKeithley(area=cfg["cell_area"])
    runner = ContinuousISOSRunner(fake, cfg, {}, threading.Event())

    # Simulate a crash: the process dies after streaming stress rows but before
    # _cleanup folds the CSV in — patch consolidation to a no-op so the
    # in-progress CSV is left behind exactly as a real crash would.
    with patch("app.instrument.keithley.time.sleep"), \
         patch("app.measurement.isos_continuous_runner.consolidate_csv_into_xlsx"):
        runner.run()

    folder = str(tmp_path)
    assert os.path.exists(runner._csv_path), "crash should leave an orphan CSV"
    orphans = find_orphan_stress_csvs([folder])
    assert len(orphans) == 1
    csv_path, xlsx = orphans[0]
    assert os.path.samefile(xlsx, runner._tracking_path)

    # Recovery consolidates it cleanly.
    sheet = consolidate_orphan(csv_path, xlsx)
    wb = openpyxl.load_workbook(xlsx)
    assert sheet in wb.sheetnames
    wb.close()
    assert os.path.exists(csv_path + ".consumed")


@pytest.mark.windows_only
def test_gui_exposes_recovery_dialog_method():
    # Importing main_window pulls in winsound, which is Windows-only; this check
    # therefore runs on the Windows CI job and is skipped on Linux/macOS.
    from app.gui.main_window import JVMeasurementApp
    # The startup scanner and the recovery-dialog method are offered by the GUI.
    assert callable(getattr(JVMeasurementApp, "_scan_for_inprogress_isos", None))
    assert callable(getattr(JVMeasurementApp, "_show_recovery_dialog", None))
