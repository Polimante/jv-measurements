"""Full-sequence integration tests for the continuous ISOS protocols.

These drive the shared ``ContinuousISOSRunner`` end-to-end against the SCPI-shaped
``FakeKeithley`` + deterministic ``VirtualClock`` (the canonical runner test
vehicle per CLAUDE.md — the runner has no DummySMU path). They lock down the
tracking-workbook contract: the number and kind of sheets produced by a complete
run, the per-scan agreement between the Log / JV / SPO / Metrics tables, and that
the autosave CSV is consolidated faithfully into the workbook.

Per the session-7 architecture decision, stress data is streamed to a single
CSV and folded into **one** ``Stress_*`` sheet at end-of-run (not per segment).
"""

from __future__ import annotations

import csv
import os
import threading
from unittest.mock import patch

import openpyxl

from app.measurement.isos_continuous_runner import ContinuousISOSRunner
from app.io.isos_tracking import read_isos_log, read_isos_metrics

from tests._isos_runner_helpers import (
    FakeKeithley,
    VirtualClock,
    make_runner_config,
    make_v1_runner_config,
)


def _sheets(path):
    wb = openpyxl.load_workbook(path)
    names = list(wb.sheetnames)
    wb.close()
    return names


def _count(names, prefix):
    return len([s for s in names if s.startswith(prefix)])


# ─────────────────────────────────────────────────────────────────────────────
# ISOS-L-1 — light soaking, MPP load, P&O stress, SPO after each scan
# ─────────────────────────────────────────────────────────────────────────────

def _run_l1(tmp_path):
    clock = VirtualClock()
    cfg = make_runner_config(
        tmp_path,
        clock=clock,
        total_duration_h=0.05,        # 180 s
        scan_interval_min=0.6,        # 36 s  → 5 main-loop scans
        stress_log_interval_s=6.0,
        load_mode="MPP",
        baseline_n_scans=4,
        baseline_rest_s=0.0,          # instant in virtual time
        spo_duration_s=1.0,           # ~5 samples → one SPO sheet per scan
        pre_scan_relaxation_s=0.0,
    )
    fake = FakeKeithley(area=cfg["cell_area"])
    stop = threading.Event()
    runner = ContinuousISOSRunner(fake, cfg, {}, stop)
    with patch("app.instrument.keithley.time.sleep"):
        runner.run()
    return runner


def test_l1_full_sequence_sheet_inventory(tmp_path):
    runner = _run_l1(tmp_path)
    names = _sheets(runner._tracking_path)

    log_rows     = read_isos_log(runner._tracking_path)
    metrics_rows = read_isos_metrics(runner._tracking_path)
    n_scans = len(log_rows)

    assert n_scans == 5, f"expected 5 main-loop scans, got {n_scans}"
    assert _count(names, "Baseline_#") == 4
    assert _count(names, "JV_#")  == n_scans            # excludes JV_Dark_#
    assert _count(names, "JV_Dark_#") == 0              # L-1 has no dark scans
    assert _count(names, "SPO_#") == n_scans
    assert _count(names, "Stress_") == 1                # single end-of-run sheet
    assert len(metrics_rows) == n_scans


def test_l1_consolidated_stress_matches_csv(tmp_path):
    runner = _run_l1(tmp_path)

    consumed = runner._csv_path + ".consumed"
    assert os.path.exists(consumed), "autosave CSV should be marked .consumed"
    assert not os.path.exists(runner._csv_path), "raw CSV should be renamed"

    with open(consumed, "r", newline="", encoding="utf-8") as fh:
        csv_rows = list(csv.reader(fh))
    assert len(csv_rows) > 1, "stress CSV should have a header + data rows"

    wb = openpyxl.load_workbook(runner._tracking_path)
    stress_name = next(s for s in wb.sheetnames if s.startswith("Stress_"))
    ws = wb[stress_name]
    sheet_rows = [[("" if c is None else str(c)) for c in row]
                  for row in ws.iter_rows(values_only=True)]
    wb.close()

    assert len(sheet_rows) == len(csv_rows)
    assert sheet_rows[0] == csv_rows[0]                 # header
    assert sheet_rows[1] == csv_rows[1]                 # first streamed sample


# ─────────────────────────────────────────────────────────────────────────────
# ISOS-V-1 — dark bias stress, tiered schedule, periodic dark + illuminated J-V
# ─────────────────────────────────────────────────────────────────────────────

def _run_v1(tmp_path):
    cfg = make_v1_runner_config(
        tmp_path,
        total_duration_h=5.0 / 60.0,          # 300 s
        tiered_schedule=[[0.0, 1.0, 1.0]],    # every 60 s (1 min) for the run
        stress_log_interval_s=6.0,
        spo_duration_s=1.0,                   # one SPO sheet per scan
    )
    fake = FakeKeithley(area=cfg["cell_area"])
    fake.set_light(False)                     # V-1 starts in the dark

    def _on_shutter(action):
        fake.set_light(action == "open")

    stop = threading.Event()
    runner = ContinuousISOSRunner(fake, cfg, {"on_shutter": _on_shutter}, stop)
    with patch("app.instrument.keithley.time.sleep"):
        runner.run()
    return runner


def test_v1_full_sequence_sheet_inventory(tmp_path):
    runner = _run_v1(tmp_path)
    names = _sheets(runner._tracking_path)

    log_rows     = read_isos_log(runner._tracking_path)
    metrics_rows = read_isos_metrics(runner._tracking_path)
    n_scans = len(log_rows)

    assert n_scans == 5, f"expected 5 tiered scans, got {n_scans}"
    assert _count(names, "JV_Dark_#") == n_scans
    assert _count(names, "JV_#") == n_scans             # JV_# excludes JV_Dark_#
    assert _count(names, "SPO_#") == n_scans
    assert _count(names, "Stress_") == 1
    assert len(metrics_rows) >= 1


def test_v1_dark_params_populated(tmp_path):
    runner = _run_v1(tmp_path)
    rows = read_isos_log(runner._tracking_path)
    assert rows
    for r in rows:
        for key in ("n_dark_rev", "J_s_dark_rev", "R_s_dark_rev",
                    "R_sh_dark_rev", "HI_dark"):
            assert r.get(key) not in (None, ""), f"{key} missing in row {r['#']}"
