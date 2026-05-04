"""
Tests for app/io/isos_exporter.py

Strategy: create synthetic tracking files in tmp_path, call the public API,
then read back with openpyxl and verify values round-trip correctly.
No instrument or GUI required.
"""

import numpy as np
import openpyxl
import pytest

from app.io.isos_exporter import (
    create_isos_d1_tracking_file,
    append_isos_d1_measurement,
    read_isos_d1_config,
    read_isos_d1_log,
    read_isos_d1_jv_curves,
    _LOG_SHEET,
    _CFG_SHEET,
    _LOG_HEADERS,
    _CFG_KEYS,
)


# ── helpers ───────────────────────────────────────────────────────────────────

_SAMPLE_CONFIG = {
    "sample_name":            "TestCell",
    "cell_area":              0.07,
    "detector_power":         110.0,
    "detector_area":          1.13,
    "start_voltage":          -0.1,
    "end_voltage":             1.1,
    "step_voltage":            0.025,
    "delay_ms":               100.0,
    "nplc":                    1.0,
    "four_wire":               False,
    "prebias_enabled":         True,
    "prebias_voltage":         1.2,
    "prebias_duration":        2.0,
    "light_soaking_enabled":   False,
    "light_soaking_duration":  300.0,
    "dark_jv":                 False,
    "electrical_load":         "OC",
    "scan_rate_mode":          "advanced",
    "target_scan_rate":        100.0,
}

_V = list(np.linspace(-0.1, 1.1, 25))
_J_F = [20.0 - 20.0 * max(0.0, v / 1.05) ** 2 for v in _V]
_J_R = [21.0 - 21.0 * max(0.0, v / 1.08) ** 2 for v in _V]


def _make_meas(pce_fwd=14.0, pce_rev=15.5, **kw):
    base = {
        "v_fwd": _V,  "j_fwd": _J_F,
        "v_rev": _V,  "j_rev": _J_R,
        "Voc_fwd": 1.02, "Jsc_fwd": 19.8, "FF_fwd": 68.5, "PCE_fwd": pce_fwd,
        "Rs_fwd":   4.2, "Rsh_fwd": 450.0,
        "Voc_rev": 1.05, "Jsc_rev": 20.5, "FF_rev": 72.1, "PCE_rev": pce_rev,
        "Rs_rev":   3.8, "Rsh_rev": 480.0,
        "HI": 10.0, "illumination": 100.0, "scan_rate": 150.0,
        # optional JV-sheet fields
        "cell_area": 0.07, "delay_ms": 100.0, "nplc": 1.0, "four_wire": False,
        "prebias_enabled": True, "prebias_voltage": 1.2, "prebias_duration": 2.0,
        "instrument_idn": "TEST IDN",
    }
    base.update(kw)
    return base


def _create(tmp_path, config=None):
    return create_isos_d1_tracking_file(str(tmp_path), config or _SAMPLE_CONFIG)


# ─────────────────────────────────────────────────────────────────────────────
# create_isos_d1_tracking_file
# ─────────────────────────────────────────────────────────────────────────────

class TestCreateTrackingFile:

    def test_returns_xlsx_path(self, tmp_path):
        path = _create(tmp_path)
        assert path.endswith(".xlsx")

    def test_file_exists(self, tmp_path):
        from pathlib import Path
        path = _create(tmp_path)
        assert Path(path).exists()

    def test_sample_name_in_filename(self, tmp_path):
        path = _create(tmp_path)
        assert "TestCell" in path

    def test_has_config_sheet(self, tmp_path):
        path = _create(tmp_path)
        wb = openpyxl.load_workbook(path)
        assert _CFG_SHEET in wb.sheetnames
        wb.close()

    def test_has_log_sheet(self, tmp_path):
        path = _create(tmp_path)
        wb = openpyxl.load_workbook(path)
        assert _LOG_SHEET in wb.sheetnames
        wb.close()

    def test_config_headers_in_row1(self, tmp_path):
        path = _create(tmp_path)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_CFG_SHEET]
        headers = [ws.cell(1, c).value for c in range(1, len(_CFG_KEYS) + 1)]
        assert headers == _CFG_KEYS
        wb.close()

    def test_t0_timestamp_written(self, tmp_path):
        path = _create(tmp_path)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_CFG_SHEET]
        t0 = None
        for c in range(1, len(_CFG_KEYS) + 1):
            if ws.cell(1, c).value == "t0_timestamp":
                t0 = ws.cell(2, c).value
                break
        wb.close()
        assert t0 is not None and len(str(t0)) > 0

    def test_log_headers_in_row1(self, tmp_path):
        path = _create(tmp_path)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_LOG_SHEET]
        headers = [ws.cell(1, c).value for c in range(1, len(_LOG_HEADERS) + 1)]
        assert headers == _LOG_HEADERS
        wb.close()

    def test_log_sheet_empty_after_creation(self, tmp_path):
        path = _create(tmp_path)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_LOG_SHEET]
        assert ws.max_row == 1  # only header row
        wb.close()

    def test_fallback_sample_name(self, tmp_path):
        cfg = dict(_SAMPLE_CONFIG, sample_name=None)
        path = create_isos_d1_tracking_file(str(tmp_path), cfg)
        assert "Sample" in path


# ─────────────────────────────────────────────────────────────────────────────
# append_isos_d1_measurement
# ─────────────────────────────────────────────────────────────────────────────

class TestAppendMeasurement:

    def test_first_append_adds_row(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        wb = openpyxl.load_workbook(path, data_only=True)
        assert wb[_LOG_SHEET].max_row == 2  # header + 1 data row
        wb.close()

    def test_sequential_appends(self, tmp_path):
        path = _create(tmp_path)
        for _ in range(3):
            append_isos_d1_measurement(path, _make_meas())
        wb = openpyxl.load_workbook(path, data_only=True)
        assert wb[_LOG_SHEET].max_row == 4  # header + 3 rows
        wb.close()

    def test_measurement_number_increments(self, tmp_path):
        path = _create(tmp_path)
        for _ in range(3):
            append_isos_d1_measurement(path, _make_meas())
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_LOG_SHEET]
        nums = [ws.cell(r, 1).value for r in range(2, 5)]
        assert nums == [1, 2, 3]
        wb.close()

    def test_pce_rev_stored_correctly(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas(pce_rev=16.2))
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_LOG_SHEET]
        # PCE_rev (%) is the 13th column (1-based)
        val = ws.cell(2, 13).value
        wb.close()
        assert val == pytest.approx(16.2, abs=0.01)

    def test_pce_fwd_stored_correctly(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas(pce_fwd=13.5))
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_LOG_SHEET]
        # PCE_fwd (%) is the 8th column (1-based)
        val = ws.cell(2, 8).value
        wb.close()
        assert val == pytest.approx(13.5, abs=0.01)

    def test_first_measurement_pce_norm_is_1(self, tmp_path):
        """First measurement must have PCE_norm = 1.0 (ratio convention, not percent)."""
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas(pce_fwd=14.0, pce_rev=15.5))
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_LOG_SHEET]
        fwd_norm = ws.cell(2, 9).value   # PCE_fwd_norm
        rev_norm = ws.cell(2, 14).value  # PCE_rev_norm
        wb.close()
        assert fwd_norm == pytest.approx(1.0, abs=0.0001)
        assert rev_norm == pytest.approx(1.0, abs=0.0001)

    def test_second_measurement_pce_norm_relative_to_first(self, tmp_path):
        """PCE_norm is PCE / PCE_t0 as a ratio (e.g. 0.9032, not 90.32)."""
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas(pce_rev=15.5))
        append_isos_d1_measurement(path, _make_meas(pce_rev=14.0))
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_LOG_SHEET]
        rev_norm2 = ws.cell(3, 14).value  # PCE_rev_norm of 2nd measurement
        wb.close()
        expected = round(14.0 / 15.5, 4)
        assert rev_norm2 == pytest.approx(expected, abs=0.001)

    def test_delta_h_zero_on_first(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_LOG_SHEET]
        delta_h = ws.cell(2, 4).value
        wb.close()
        assert delta_h == pytest.approx(0.0, abs=0.001)

    def test_hi_stored_correctly(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas(HI=9.5))
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_LOG_SHEET]
        val = ws.cell(2, 15).value
        wb.close()
        assert val == pytest.approx(9.5, abs=0.01)

    def test_jv_sheet_created_per_measurement(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        wb = openpyxl.load_workbook(path)
        jv_sheets = [s for s in wb.sheetnames if s.startswith("JV_")]
        wb.close()
        # One combined sheet per measurement (not separate Fwd/Rev sheets)
        assert len(jv_sheets) == 1

    def test_jv_sheet_name_format(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        wb = openpyxl.load_workbook(path)
        jv_sheets = [s for s in wb.sheetnames if s.startswith("JV_")]
        wb.close()
        # Format: JV_#001_tXXXXh
        assert jv_sheets[0].startswith("JV_#001_t")

    def test_jv_sheet_has_header_and_scalar_and_data_rows(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        wb = openpyxl.load_workbook(path, data_only=True)
        jv_sheet = next(n for n in wb.sheetnames if n.startswith("JV_"))
        ws = wb[jv_sheet]
        # row 1 = headers, row 2 = scalars, rows 3+ = data
        assert ws.max_row == 2 + len(_V)
        wb.close()

    def test_jv_sheet_has_six_data_columns(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        wb = openpyxl.load_workbook(path, data_only=True)
        jv_sheet = next(n for n in wb.sheetnames if n.startswith("JV_"))
        ws = wb[jv_sheet]
        # Cols A-F: V_fwd, J_fwd, P_fwd, V_rev, J_rev, P_rev
        row3 = [ws.cell(3, c).value for c in range(1, 7)]
        wb.close()
        assert all(v is not None for v in row3)

    def test_notes_stored(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas(notes="test note"))
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_LOG_SHEET]
        notes_col = len(_LOG_HEADERS)
        val = ws.cell(2, notes_col).value
        wb.close()
        assert val == "test note"

    def test_notes_defaults_to_empty(self, tmp_path):
        path = _create(tmp_path)
        meas = _make_meas()
        meas.pop("notes", None)
        append_isos_d1_measurement(path, meas)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[_LOG_SHEET]
        val = ws.cell(2, len(_LOG_HEADERS)).value
        wb.close()
        assert val in (None, "")


# ─────────────────────────────────────────────────────────────────────────────
# read_isos_d1_config
# ─────────────────────────────────────────────────────────────────────────────

class TestReadConfig:

    def test_reads_sample_name(self, tmp_path):
        path = _create(tmp_path)
        cfg = read_isos_d1_config(path)
        assert cfg["sample_name"] == "TestCell"

    def test_reads_cell_area(self, tmp_path):
        path = _create(tmp_path)
        cfg = read_isos_d1_config(path)
        assert cfg["cell_area"] == pytest.approx(0.07, abs=0.001)

    def test_reads_t0_timestamp(self, tmp_path):
        path = _create(tmp_path)
        cfg = read_isos_d1_config(path)
        assert "t0_timestamp" in cfg
        assert cfg["t0_timestamp"] is not None

    def test_all_cfg_keys_present(self, tmp_path):
        path = _create(tmp_path)
        cfg = read_isos_d1_config(path)
        for key in _CFG_KEYS:
            assert key in cfg

    def test_nonexistent_file_returns_empty(self):
        cfg = read_isos_d1_config("/nonexistent/path/file.xlsx")
        assert cfg == {}

    def test_wrong_sheet_returns_empty(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.active.title = "WrongSheet"
        p = str(tmp_path / "bad.xlsx")
        wb.save(p)
        wb.close()
        assert read_isos_d1_config(p) == {}


# ─────────────────────────────────────────────────────────────────────────────
# read_isos_d1_log
# ─────────────────────────────────────────────────────────────────────────────

class TestReadLog:

    def test_empty_log_returns_empty_list(self, tmp_path):
        path = _create(tmp_path)
        rows = read_isos_d1_log(path)
        assert rows == []

    def test_one_measurement_returns_one_row(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        rows = read_isos_d1_log(path)
        assert len(rows) == 1

    def test_three_measurements_returns_three_rows(self, tmp_path):
        path = _create(tmp_path)
        for _ in range(3):
            append_isos_d1_measurement(path, _make_meas())
        rows = read_isos_d1_log(path)
        assert len(rows) == 3

    def test_row_has_expected_keys(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        row = read_isos_d1_log(path)[0]
        for key in _LOG_HEADERS:
            assert key in row

    def test_pce_rev_roundtrip(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas(pce_rev=16.2))
        row = read_isos_d1_log(path)[0]
        assert row["PCE_rev (%)"] == pytest.approx(16.2, abs=0.01)

    def test_hi_roundtrip(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas(HI=8.7))
        row = read_isos_d1_log(path)[0]
        assert row["HI (%)"] == pytest.approx(8.7, abs=0.01)

    def test_number_column_is_sequential(self, tmp_path):
        path = _create(tmp_path)
        for _ in range(3):
            append_isos_d1_measurement(path, _make_meas())
        rows = read_isos_d1_log(path)
        assert [r["#"] for r in rows] == [1, 2, 3]

    def test_nonexistent_file_returns_empty(self):
        assert read_isos_d1_log("/nonexistent/file.xlsx") == []

    def test_wrong_sheet_returns_empty(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.active.title = "WrongSheet"
        p = str(tmp_path / "bad.xlsx")
        wb.save(p)
        wb.close()
        assert read_isos_d1_log(p) == []


# ─────────────────────────────────────────────────────────────────────────────
# read_isos_d1_jv_curves
# ─────────────────────────────────────────────────────────────────────────────

class TestReadJVCurves:

    def test_no_curves_before_append(self, tmp_path):
        path = _create(tmp_path)
        curves = read_isos_d1_jv_curves(path)
        assert curves == []

    def test_one_append_gives_two_curves(self, tmp_path):
        # One combined sheet → 2 entries (Fwd + Rev)
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        curves = read_isos_d1_jv_curves(path)
        assert len(curves) == 2

    def test_two_appends_give_four_curves(self, tmp_path):
        # Two sheets → 4 entries
        path = _create(tmp_path)
        for _ in range(2):
            append_isos_d1_measurement(path, _make_meas())
        curves = read_isos_d1_jv_curves(path)
        assert len(curves) == 4

    def test_directions_are_fwd_and_rev(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        curves = read_isos_d1_jv_curves(path)
        directions = {c["direction"] for c in curves}
        assert directions == {"Fwd", "Rev"}

    def test_curve_has_v_and_j_lists(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        curves = read_isos_d1_jv_curves(path)
        for c in curves:
            assert "v" in c and "j" in c
            assert len(c["v"]) == len(_V)
            assert len(c["j"]) == len(_V)

    def test_voltage_values_roundtrip(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        curves = read_isos_d1_jv_curves(path)
        fwd = next(c for c in curves if c["direction"] == "Fwd")
        assert fwd["v"][0] == pytest.approx(_V[0], abs=0.001)
        assert fwd["v"][-1] == pytest.approx(_V[-1], abs=0.001)

    def test_curve_elapsed_h_is_non_negative(self, tmp_path):
        path = _create(tmp_path)
        append_isos_d1_measurement(path, _make_meas())
        curves = read_isos_d1_jv_curves(path)
        for c in curves:
            assert c["elapsed_h"] >= 0

    def test_nonexistent_file_returns_empty(self):
        assert read_isos_d1_jv_curves("/nonexistent/file.xlsx") == []
