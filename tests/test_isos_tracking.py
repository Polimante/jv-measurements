"""
Tests for app/io/isos_tracking.py (generalized ISOS tracking module).

Strategy: create synthetic tracking files in tmp_path, call both the generic
and legacy D-1 public APIs, then read back and verify values round-trip. Also
verifies that a hand-built legacy (ISOS_D1_* sheet) file still opens. No
instrument or GUI required.
"""

import numpy as np
import openpyxl
import pytest

from app.io.isos_tracking import (
    create_isos_tracking_file,
    append_isos_measurement,
    read_isos_config,
    read_isos_log,
    read_isos_jv_curves,
    create_isos_d1_tracking_file,
    append_isos_d1_measurement,
    read_isos_d1_config,
    read_isos_d1_log,
    read_isos_d1_jv_curves,
    _LOG_SHEET,
    _CFG_SHEET,
    _LEGACY_LOG_SHEET,
    _LEGACY_CFG_SHEET,
    _LOG_HEADERS,
    _LOG_HEADERS_CORE,
    _CFG_KEYS,
)


# ── helpers ───────────────────────────────────────────────────────────────────

_SAMPLE_CONFIG = {
    "sample_name":            "TestCell",
    "cell_area":              0.07,
    "detector_power":         110.0,
    "detector_area":          1.13,
    "start_voltage":          -0.1,
    "end_voltage":             1.1,
    "step_voltage":            0.025,
    "delay_ms":               100.0,
    "nplc":                    1.0,
    "four_wire":               False,
    "prebias_enabled":         True,
    "prebias_voltage":         1.2,
    "prebias_duration":        2.0,
    "light_soaking_enabled":   False,
    "light_soaking_duration":  300.0,
    "dark_jv":                 False,
    "electrical_load":         "OC",
    "scan_rate_mode":          "advanced",
    "target_scan_rate":        100.0,
}

_V = list(np.linspace(-0.1, 1.1, 25))
_J_F = [20.0 - 20.0 * max(0.0, v / 1.05) ** 2 for v in _V]
_J_R = [21.0 - 21.0 * max(0.0, v / 1.08) ** 2 for v in _V]


def _make_meas(pce_fwd=14.0, pce_rev=15.5, **kw):
    base = {
        "v_fwd": _V,  "j_fwd": _J_F,
        "v_rev": _V,  "j_rev": _J_R,
        "Voc_fwd": 1.02, "Jsc_fwd": 19.8, "FF_fwd": 68.5, "PCE_fwd": pce_fwd,
        "Rs_fwd":   4.2, "Rsh_fwd": 450.0,
        "Voc_rev": 1.05, "Jsc_rev": 20.5, "FF_rev": 72.1, "PCE_rev": pce_rev,
        "Rs_rev":   3.8, "Rsh_rev": 480.0,
        "HI": 10.0, "illumination": 100.0, "scan_rate": 150.0,
        "cell_area": 0.07, "delay_ms": 100.0, "nplc": 1.0, "four_wire": False,
        "prebias_enabled": True, "prebias_voltage": 1.2, "prebias_duration": 2.0,
        "instrument_idn": "TEST IDN",
    }
    base.update(kw)
    return base


def _create(tmp_path, config=None, protocol_id="ISOS-D-1"):
    cfg = dict(config or _SAMPLE_CONFIG)
    cfg["protocol_id"] = protocol_id
    return create_isos_tracking_file(str(tmp_path), cfg)


def _make_legacy_d1_file(tmp_path):
    """Build an old-format D-1 file with ISOS_D1_* sheets and the original
    22-column log layout, populated with one measurement row."""
    legacy_cfg_keys = [
        "sample_name", "cell_area", "detector_power", "detector_area",
        "start_voltage", "end_voltage", "step_voltage",
        "delay_ms", "nplc", "four_wire",
        "prebias_enabled", "prebias_voltage", "prebias_duration",
        "light_soaking_enabled", "light_soaking_duration",
        "dark_jv", "electrical_load",
        "scan_rate_mode", "target_scan_rate",
        "t0_timestamp",
    ]
    wb = openpyxl.Workbook()
    ws_cfg = wb.active
    ws_cfg.title = _LEGACY_CFG_SHEET
    for c, k in enumerate(legacy_cfg_keys, 1):
        ws_cfg.cell(1, c, k)
    cfg_vals = [
        "LegacyCell", 0.09, 100.0, 1.0, -0.2, 1.2, 0.02,
        50.0, 1.0, False, True, 1.1, 3.0, False, 0.0, False, "OC",
        "advanced", 100.0, "2024-01-01 00:00:00",
    ]
    ws_cfg.append(cfg_vals)

    ws_log = wb.create_sheet(_LEGACY_LOG_SHEET)
    for c, h in enumerate(_LOG_HEADERS_CORE, 1):
        ws_log.cell(1, c, h)
    # one measurement row (22 columns)
    ws_log.append([
        1, "2024-01-01 00:00:00", 0.0, 0.0,
        1.01, 19.0, 70.0, 13.3, 1.0,
        1.04, 20.0, 73.0, 15.1, 1.0,
        8.0,
        4.0, 400.0, 3.5, 420.0,
        100.0, 120.0, "legacy note",
    ])

    # one legacy JV sheet (same column format)
    ws_jv = wb.create_sheet("JV_#001_t0000h")
    ws_jv.cell(1, 1, "Forward Voltage (V)")
    ws_jv.append([None] * 6)  # scalar row placeholder
    for i in range(len(_V)):
        ws_jv.append([_V[i], _J_F[i], _V[i] * _J_F[i],
                      _V[i], _J_R[i], _V[i] * _J_R[i]])

    p = str(tmp_path / "LegacyCell_ISOS-D1_legacy.xlsx")
    wb.save(p)
    wb.close()
    return p


# ─────────────────────────────────────────────────────────────────────────────
# create_isos_tracking_file
# ─────────────────────────────────────────────────────────────────────────────

class TestCreateTrackingFile:

    def test_returns_xlsx_path(self, tmp_path):
        assert _create(tmp_path).endswith(".xlsx")

    def test_file_exists(self, tmp_path):
        from pathlib import Path
        assert Path(_create(tmp_path)).exists()

    def test_sample_name_in_filename(self, tmp_path):
        assert "TestCell" in _create(tmp_path)

    def test_has_generic_sheets(self, tmp_path):
        wb = openpyxl.load_workbook(_create(tmp_path))
        assert _CFG_SHEET in wb.sheetnames
        assert _LOG_SHEET in wb.sheetnames
        wb.close()

    def test_protocol_id_is_first_config_column(self, tmp_path):
        wb = openpyxl.load_workbook(_create(tmp_path), data_only=True)
        ws = wb[_CFG_SHEET]
        assert ws.cell(1, 1).value == "protocol_id"
        assert ws.cell(2, 1).value == "ISOS-D-1"
        wb.close()

    def test_config_headers_in_row1(self, tmp_path):
        wb = openpyxl.load_workbook(_create(tmp_path), data_only=True)
        ws = wb[_CFG_SHEET]
        headers = [ws.cell(1, c).value for c in range(1, len(_CFG_KEYS) + 1)]
        assert headers == _CFG_KEYS
        wb.close()

    def test_log_headers_in_row1(self, tmp_path):
        wb = openpyxl.load_workbook(_create(tmp_path), data_only=True)
        ws = wb[_LOG_SHEET]
        headers = [ws.cell(1, c).value for c in range(1, len(_LOG_HEADERS) + 1)]
        assert headers == _LOG_HEADERS
        wb.close()

    def test_t0_timestamp_written(self, tmp_path):
        cfg = read_isos_config(_create(tmp_path))
        assert cfg.get("t0_timestamp")

    def test_filename_label_per_protocol(self, tmp_path):
        assert "ISOS-D1" in _create(tmp_path, protocol_id="ISOS-D-1")
        assert "ISOS-L1" in _create(tmp_path, protocol_id="ISOS-L-1")
        assert "ISOS-V1" in _create(tmp_path, protocol_id="ISOS-V-1")


# ─────────────────────────────────────────────────────────────────────────────
# Generic round-trip across protocol ids
# ─────────────────────────────────────────────────────────────────────────────

class TestGenericRoundTrip:

    @pytest.mark.parametrize("protocol_id", ["ISOS-D-1", "ISOS-L-1", "ISOS-V-1"])
    def test_config_protocol_id_roundtrips(self, tmp_path, protocol_id):
        path = _create(tmp_path, protocol_id=protocol_id)
        cfg = read_isos_config(path)
        assert cfg["protocol_id"] == protocol_id

    @pytest.mark.parametrize("protocol_id", ["ISOS-D-1", "ISOS-L-1", "ISOS-V-1"])
    def test_append_and_read_log(self, tmp_path, protocol_id):
        path = _create(tmp_path, protocol_id=protocol_id)
        append_isos_measurement(path, _make_meas(pce_rev=16.2))
        rows = read_isos_log(path)
        assert len(rows) == 1
        assert rows[0]["PCE_rev (%)"] == pytest.approx(16.2, abs=0.01)
        assert rows[0]["#"] == 1

    @pytest.mark.parametrize("protocol_id", ["ISOS-D-1", "ISOS-L-1", "ISOS-V-1"])
    def test_append_and_read_curves(self, tmp_path, protocol_id):
        path = _create(tmp_path, protocol_id=protocol_id)
        append_isos_measurement(path, _make_meas())
        curves = read_isos_jv_curves(path)
        assert {c["direction"] for c in curves} == {"Fwd", "Rev"}
        for c in curves:
            assert len(c["v"]) == len(_V)

    def test_new_nullable_columns_default_none(self, tmp_path):
        path = _create(tmp_path)
        append_isos_measurement(path, _make_meas())
        row = read_isos_log(path)[0]
        assert "Load" in row and row["Load"] is None
        assert "PCE_SPO" in row and row["PCE_SPO"] is None
        assert "failure_reason" in row and row["failure_reason"] is None

    def test_new_nullable_columns_persist_when_supplied(self, tmp_path):
        path = _create(tmp_path)
        append_isos_measurement(path, _make_meas(
            load="MPP", setpoint_v=0.82, mpp_engine="PO",
            pce_spo=14.7, failure_detected=True, failure_reason="Voc collapse",
        ))
        row = read_isos_log(path)[0]
        assert row["Load"] == "MPP"
        assert row["Setpoint_V"] == pytest.approx(0.82, abs=1e-6)
        assert row["MPP_Engine"] == "PO"
        assert row["PCE_SPO"] == pytest.approx(14.7, abs=1e-6)
        assert row["failure_reason"] == "Voc collapse"

    def test_pce_norm_first_is_one(self, tmp_path):
        path = _create(tmp_path)
        append_isos_measurement(path, _make_meas(pce_fwd=14.0, pce_rev=15.5))
        row = read_isos_log(path)[0]
        assert row["PCE_fwd_norm"] == pytest.approx(1.0, abs=1e-4)
        assert row["PCE_rev_norm"] == pytest.approx(1.0, abs=1e-4)

    def test_pce_norm_relative_to_first(self, tmp_path):
        path = _create(tmp_path)
        append_isos_measurement(path, _make_meas(pce_rev=15.5))
        append_isos_measurement(path, _make_meas(pce_rev=14.0))
        rows = read_isos_log(path)
        assert rows[1]["PCE_rev_norm"] == pytest.approx(round(14.0 / 15.5, 4), abs=1e-3)


# ─────────────────────────────────────────────────────────────────────────────
# Legacy D-1 wrapper functions
# ─────────────────────────────────────────────────────────────────────────────

class TestD1Wrappers:

    def test_create_wrapper_sets_protocol_id(self, tmp_path):
        path = create_isos_d1_tracking_file(str(tmp_path), _SAMPLE_CONFIG)
        cfg = read_isos_d1_config(path)
        assert cfg["protocol_id"] == "ISOS-D-1"

    def test_wrapper_full_roundtrip(self, tmp_path):
        path = create_isos_d1_tracking_file(str(tmp_path), _SAMPLE_CONFIG)
        for _ in range(3):
            append_isos_d1_measurement(path, _make_meas())
        rows = read_isos_d1_log(path)
        assert [r["#"] for r in rows] == [1, 2, 3]
        curves = read_isos_d1_jv_curves(path)
        assert len(curves) == 6  # 3 measurements × (Fwd + Rev)

    def test_wrapper_config_reads_sample_name(self, tmp_path):
        path = create_isos_d1_tracking_file(str(tmp_path), _SAMPLE_CONFIG)
        assert read_isos_d1_config(path)["sample_name"] == "TestCell"

    def test_fallback_sample_name(self, tmp_path):
        cfg = dict(_SAMPLE_CONFIG, sample_name=None)
        assert "Sample" in create_isos_d1_tracking_file(str(tmp_path), cfg)


# ─────────────────────────────────────────────────────────────────────────────
# Legacy (old-format) file compatibility
# ─────────────────────────────────────────────────────────────────────────────

class TestLegacyFileCompat:

    def test_legacy_config_opens(self, tmp_path):
        path = _make_legacy_d1_file(tmp_path)
        cfg = read_isos_config(path)
        assert cfg["sample_name"] == "LegacyCell"
        assert cfg["t0_timestamp"] == "2024-01-01 00:00:00"

    def test_legacy_log_opens(self, tmp_path):
        path = _make_legacy_d1_file(tmp_path)
        rows = read_isos_log(path)
        assert len(rows) == 1
        assert rows[0]["#"] == 1
        assert rows[0]["PCE_rev (%)"] == pytest.approx(15.1, abs=0.01)
        assert rows[0]["Notes"] == "legacy note"

    def test_legacy_log_has_no_phantom_keys(self, tmp_path):
        path = _make_legacy_d1_file(tmp_path)
        row = read_isos_log(path)[0]
        # Only the 22 core headers exist in a legacy file — no None keys.
        assert None not in row
        assert set(row.keys()) == set(_LOG_HEADERS_CORE)

    def test_legacy_curves_open(self, tmp_path):
        path = _make_legacy_d1_file(tmp_path)
        curves = read_isos_jv_curves(path)
        assert {c["direction"] for c in curves} == {"Fwd", "Rev"}

    def test_legacy_wrapper_reads(self, tmp_path):
        path = _make_legacy_d1_file(tmp_path)
        assert read_isos_d1_config(path)["sample_name"] == "LegacyCell"
        assert len(read_isos_d1_log(path)) == 1


# ─────────────────────────────────────────────────────────────────────────────
# Error handling
# ─────────────────────────────────────────────────────────────────────────────

class TestErrorHandling:

    def test_nonexistent_config_returns_empty(self):
        assert read_isos_config("/nonexistent/file.xlsx") == {}

    def test_nonexistent_log_returns_empty(self):
        assert read_isos_log("/nonexistent/file.xlsx") == []

    def test_nonexistent_curves_returns_empty(self):
        assert read_isos_jv_curves("/nonexistent/file.xlsx") == []

    def test_wrong_sheet_returns_empty(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.active.title = "WrongSheet"
        p = str(tmp_path / "bad.xlsx")
        wb.save(p)
        wb.close()
        assert read_isos_config(p) == {}
        assert read_isos_log(p) == []
