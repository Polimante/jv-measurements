"""Tests for save_light_intensity_jv_consolidated() in light_intensity_exporter.py.

Strategy: build synthetic IntensityResult lists, call the exporter, then reopen
the workbook with openpyxl and verify every relevant sheet and cell.

No instrument, GUI, or measurement thread required.
"""

from __future__ import annotations

import math
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from app.io.light_intensity_exporter import save_light_intensity_jv_consolidated
from app.measurement.light_intensity_jv import IntensityResult


# ── Helpers ────────────────────────────────────────────────────────────────────

_SCAN_PARAMS = {
    "start_v":   0.0,
    "end_v":     1.1,
    "step_v":    0.05,
    "delay_ms":  100.0,
    "nplc":      1.0,
    "cell_area": 0.09,
    "four_wire": False,
    "reference_illumination_mw_cm2": 100.0,
}

_INTENSITIES = [0.25, 0.5, 0.75, 1.0]


def _make_v_j(n: int = 23):
    v = np.linspace(0.0, 1.1, n)
    j = 20.0 - 20.0 * np.exp(v / 0.025) / np.exp(1.1 / 0.025)
    j = np.clip(j, -20, 22)
    return v, j


def _make_result(idx: int, intensity: float, status: str = "completed") -> IntensityResult:
    v, j = _make_v_j()
    p     = v * j
    r     = IntensityResult(
        index=idx,
        intensity=intensity,
        status=status,
    )
    if status == "completed":
        r.v_fwd = v;  r.j_fwd = j
        r.v_rev = v;  r.j_rev = j
        illum_mw = intensity * 100.0
        r.Voc_fwd = 1.05 + 0.02 * idx
        r.Jsc_fwd = 18.0 + idx
        r.FF_fwd  = 72.0
        r.PCE_fwd = r.Voc_fwd * r.Jsc_fwd * r.FF_fwd / illum_mw
        r.Rs_fwd  = 5.0
        r.Rsh_fwd = 500.0
        r.Voc_rev = r.Voc_fwd + 0.005
        r.Jsc_rev = r.Jsc_fwd
        r.FF_rev  = 73.0
        r.PCE_rev = r.Voc_rev * r.Jsc_rev * r.FF_rev / illum_mw
        r.Rs_rev  = 5.2
        r.Rsh_rev = 490.0
        r.HI      = (r.PCE_rev - r.PCE_fwd) / r.PCE_rev * 100.0
    return r


def _run_export(tmp_path: Path, results=None, intensities=None,
                aborted_at=None, auto_fit=False):
    if intensities is None:
        intensities = _INTENSITIES
    if results is None:
        results = [_make_result(i, L) for i, L in enumerate(intensities)]
    return save_light_intensity_jv_consolidated(
        output_folder=tmp_path,
        sample_name="ExportTest",
        intensities=intensities,
        results=results,
        scan_params=_SCAN_PARAMS,
        stabilization_s=5.0,
        auto_fit=auto_fit,
        aborted_at=aborted_at,
        instrument_idn="KEITHLEY,2400",
    )


# ── File creation ──────────────────────────────────────────────────────────────

class TestFileCreation:
    def test_returns_path_that_exists(self, tmp_path):
        path = _run_export(tmp_path)
        assert path.exists()

    def test_filename_contains_sample_name(self, tmp_path):
        path = _run_export(tmp_path)
        assert "ExportTest" in path.name

    def test_filename_contains_LightIntensityJV(self, tmp_path):
        path = _run_export(tmp_path)
        assert "LightIntensityJV" in path.name

    def test_file_is_valid_xlsx(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        assert wb is not None
        wb.close()


# ── Config sheet ───────────────────────────────────────────────────────────────

class TestConfigSheet:
    def _get_config_kv(self, path):
        wb = openpyxl.load_workbook(str(path), data_only=True)
        ws = wb["Config"]
        kv = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] is not None:
                kv[str(row[0])] = row[1]
        wb.close()
        return kv

    def test_config_sheet_present(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        assert "Config" in wb.sheetnames
        wb.close()

    def test_sample_name_stored(self, tmp_path):
        kv = self._get_config_kv(_run_export(tmp_path))
        assert kv.get("Sample Name") == "ExportTest"

    def test_total_intensities_planned(self, tmp_path):
        kv = self._get_config_kv(_run_export(tmp_path))
        assert int(kv["Total Intensities Planned"]) == len(_INTENSITIES)

    def test_instrument_idn_stored(self, tmp_path):
        kv = self._get_config_kv(_run_export(tmp_path))
        assert kv.get("Instrument IDN") == "KEITHLEY,2400"

    def test_cell_area_stored(self, tmp_path):
        kv = self._get_config_kv(_run_export(tmp_path))
        assert abs(float(kv["Cell Area (cm²)"]) - 0.09) < 1e-9

    def test_abort_note_present_when_aborted(self, tmp_path):
        path = _run_export(tmp_path, aborted_at=2)
        kv   = self._get_config_kv(path)
        assert "Note" in kv
        assert kv["Note"] is not None


# ── Summary sheet ──────────────────────────────────────────────────────────────

class TestSummarySheet:
    def _read_summary(self, path):
        wb  = openpyxl.load_workbook(str(path), data_only=True)
        ws  = wb["Summary"]
        rows = [
            row for row in ws.iter_rows(min_row=2, values_only=True)
            if row and row[0] is not None
        ]
        wb.close()
        return rows

    def test_summary_sheet_present(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        assert "Summary" in wb.sheetnames
        wb.close()

    def test_one_row_per_intensity(self, tmp_path):
        rows = self._read_summary(_run_export(tmp_path))
        assert len(rows) == len(_INTENSITIES)

    def test_intensity_values_match(self, tmp_path):
        rows = self._read_summary(_run_export(tmp_path))
        for row, expected in zip(rows, _INTENSITIES):
            assert abs(float(row[1]) - expected) < 1e-9

    def test_pce_rev_positive(self, tmp_path):
        rows = self._read_summary(_run_export(tmp_path))
        for row in rows:
            pce = float(row[9])
            assert pce > 0, f"PCE_rev={pce}"

    def test_status_column_is_completed(self, tmp_path):
        rows = self._read_summary(_run_export(tmp_path))
        for row in rows:
            assert row[12] == "completed"

    def test_skipped_row_has_skipped_status(self, tmp_path):
        results = [_make_result(i, L) for i, L in enumerate(_INTENSITIES)]
        results[1] = _make_result(1, _INTENSITIES[1], status="skipped")
        rows = self._read_summary(_run_export(tmp_path, results=results))
        assert rows[1][12] == "skipped"


# ── LightIdealityAnalysis sheet ───────────────────────────────────────────────

class TestLightIdealitySheet:
    def test_ideality_sheet_present_for_4_intensities(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        assert "LightIdealityAnalysis" in wb.sheetnames
        wb.close()

    def test_ideality_sheet_absent_for_2_completed(self, tmp_path):
        intensities = [0.5, 1.0]
        results     = [_make_result(i, L) for i, L in enumerate(intensities)]
        path = _run_export(tmp_path, results=results, intensities=intensities)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        assert "LightIdealityAnalysis" not in wb.sheetnames
        wb.close()

    def test_n_idL_is_numeric(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        ws   = wb["LightIdealityAnalysis"]
        n_idL_fwd = ws.cell(2, 2).value
        wb.close()
        assert n_idL_fwd is not None
        assert math.isfinite(float(n_idL_fwd))

    def test_r_squared_between_0_and_1(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        ws   = wb["LightIdealityAnalysis"]
        r2   = float(ws.cell(6, 2).value or float("nan"))
        wb.close()
        assert 0.0 <= r2 <= 1.0


# ── JV sheets ──────────────────────────────────────────────────────────────────

class TestJVSheets:
    def test_jv_fwd_sheet_present(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        jv_sheets = [s for s in wb.sheetnames if s.startswith("JV_") and s.endswith("_Fwd")]
        wb.close()
        assert len(jv_sheets) == len(_INTENSITIES)

    def test_jv_rev_sheet_present(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        jv_sheets = [s for s in wb.sheetnames if s.startswith("JV_") and s.endswith("_Rev")]
        wb.close()
        assert len(jv_sheets) == len(_INTENSITIES)

    def test_jv_sheet_has_data_rows(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        fwd_sheet = next(s for s in wb.sheetnames if s.startswith("JV_") and s.endswith("_Fwd"))
        ws   = wb[fwd_sheet]
        data_rows = [
            row for row in ws.iter_rows(min_row=2, values_only=True)
            if row and row[0] is not None
        ]
        wb.close()
        assert len(data_rows) == 23   # _make_v_j uses n=23 points

    def test_jv_power_equals_v_times_j(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        fwd_sheet = next(s for s in wb.sheetnames if s.startswith("JV_") and s.endswith("_Fwd"))
        ws   = wb[fwd_sheet]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] is not None:
                v, j, p = float(row[0]), float(row[1]), float(row[2])
                assert abs(p - v * j) < 1e-10
        wb.close()

    def test_skipped_result_has_no_jv_sheet(self, tmp_path):
        results = [_make_result(i, L) for i, L in enumerate(_INTENSITIES)]
        results[2] = _make_result(2, _INTENSITIES[2], status="skipped")
        path = _run_export(tmp_path, results=results)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        tag  = f"{_INTENSITIES[2]:.4g}".replace(".", "_")
        sheet_fwd = f"JV_{tag}sun_Fwd"
        assert sheet_fwd not in wb.sheetnames
        wb.close()


# ── Sheet order ────────────────────────────────────────────────────────────────

class TestSheetOrder:
    def test_config_is_first_sheet(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        assert wb.sheetnames[0] == "Config"
        wb.close()

    def test_summary_is_second_sheet(self, tmp_path):
        path = _run_export(tmp_path)
        wb   = openpyxl.load_workbook(str(path), data_only=True)
        assert wb.sheetnames[1] == "Summary"
        wb.close()
