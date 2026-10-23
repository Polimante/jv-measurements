"""Integration tests for LightIntensityJVMeasurement orchestrator.

Runs hardware-free via DummySMU (sleep_scale=0) at multiple intensities.

Covers:
  - Full 4-intensity sequence completes: all results are "completed"
  - PV parameters are finite and physically sane after each step
  - stop_event aborts: remaining results are "skipped"
  - Consolidated file is created on disk after a successful run
  - LightIdealityAnalysis sheet present when >= 3 intensities complete
  - n_idL recoverable within reasonable bounds (0.5–3.5)
"""

from __future__ import annotations

import math
import threading
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from app.instrument.dummy_smu import DummySMU
from app.measurement.light_intensity_jv import LightIntensityJVMeasurement


# ── Shared test fixtures ───────────────────────────────────────────────────────

_SCAN_PARAMS = {
    "start_v":                     -0.05,
    "end_v":                        1.3,
    "step_v":                       0.025,
    "delay_ms":                     0.0,
    "nplc":                         0.01,
    "cell_area":                    0.09,
    "four_wire":                    False,
    "reference_illumination_mw_cm2": 100.0,
}

_INTENSITIES = [0.25, 0.5, 0.75, 1.0]


def _make_measurement(
    tmp_path: Path,
    intensities=None,
    stop_event=None,
    on_prompt_user=None,
    auto_fit=False,
    prebias=None,
    ls=None,
):
    """Return a configured LightIntensityJVMeasurement with DummySMU."""
    smu = DummySMU(sleep_scale=0.0)
    smu.connect()
    if stop_event is None:
        stop_event = threading.Event()

    def _prompt(msg, intensity):
        smu.set_illumination_factor(intensity)
        return True

    meas = LightIntensityJVMeasurement(
        sourcemeter=smu,
        intensities=intensities if intensities is not None else _INTENSITIES,
        scan_params=_SCAN_PARAMS,
        stabilization_time_s=0,
        pre_bias_settings=prebias or {"enabled": False, "voltage": 0.0, "duration": 0.0},
        light_soaking_settings=ls or {"enabled": False, "duration": 0.0},
        sample_name="TestCell",
        output_folder=tmp_path,
        auto_single_diode_fit=auto_fit,
        on_status=lambda s: None,
        on_prompt_user=on_prompt_user or _prompt,
        on_progress=lambda i, n, s: None,
        on_measurement_done=lambda i, r: None,
        on_complete=lambda p: None,
        stop_event=stop_event,
    )
    return meas, smu


# ── Full sequence ─────────────────────────────────────────────────────────────

class TestFullSequence:
    def test_all_results_completed(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        assert len(meas.results) == len(_INTENSITIES)
        for r in meas.results:
            assert r.status == "completed", (
                f"intensity={r.intensity}: status={r.status}, err={r.error_message}"
            )

    def test_pv_parameters_finite(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        for r in meas.results:
            for attr in ("Voc_rev", "Jsc_rev", "FF_rev", "PCE_rev"):
                val = getattr(r, attr)
                assert math.isfinite(val), f"{attr}={val} at intensity={r.intensity}"

    def test_voc_positive_and_below_bandgap(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        for r in meas.results:
            assert 0.0 < r.Voc_rev < 2.0, (
                f"Voc_rev={r.Voc_rev:.4f} out of range at {r.intensity} sun"
            )

    def test_jsc_positive(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        for r in meas.results:
            assert r.Jsc_rev > 0, f"Jsc_rev={r.Jsc_rev} not positive"

    def test_ff_between_0_and_100(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        for r in meas.results:
            assert 0.0 < r.FF_rev <= 100.0

    def test_voc_increases_with_intensity(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        vocs = [r.Voc_rev for r in meas.results]
        # Non-decreasing (coarse voltage step can make adjacent intensities tie)
        for i in range(len(vocs) - 1):
            assert vocs[i] <= vocs[i + 1], (
                f"Voc should be non-decreasing with intensity: {vocs}"
            )
        # At least some increase from lowest to highest intensity
        assert vocs[-1] > vocs[0], f"Voc at highest intensity should exceed lowest: {vocs}"

    def test_consolidated_file_created(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        xlsx_files = list(tmp_path.glob("*_LightIntensityJV_*.xlsx"))
        assert len(xlsx_files) == 1, f"Expected 1 consolidated file, found: {xlsx_files}"

    def test_individual_backup_files_created(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        # save_data creates files with the sample name prefix
        all_xlsx = list(tmp_path.glob("*.xlsx"))
        # Individual backups + 1 consolidated = len(intensities)+1
        assert len(all_xlsx) == len(_INTENSITIES) + 1

    def test_n_idL_in_physical_range(self, tmp_path):
        """n_idL from the fit should be between 0.5 and 3.5 for DummyCell."""
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        xlsx = next(tmp_path.glob("*_LightIntensityJV_*.xlsx"))
        wb   = openpyxl.load_workbook(str(xlsx), data_only=True)
        assert "LightIdealityAnalysis" in wb.sheetnames, (
            "LightIdealityAnalysis sheet missing"
        )
        ws = wb["LightIdealityAnalysis"]
        # Row 2 = n_idL, col B = fwd value
        n_idL_fwd = float(ws.cell(2, 2).value or float("nan"))
        wb.close()
        assert 0.5 <= n_idL_fwd <= 3.5, (
            f"n_idL_fwd={n_idL_fwd} outside expected physical range"
        )


# ── Stop event ────────────────────────────────────────────────────────────────

class TestStopEvent:
    def test_abort_after_first_marks_rest_skipped(self, tmp_path):
        """Set stop_event before the second prompt to abort mid-sequence."""
        stop_event = threading.Event()
        call_count = [0]

        def _prompt(msg, intensity):
            call_count[0] += 1
            if call_count[0] == 1:
                # Let first measurement proceed
                smu_ref[0].set_illumination_factor(intensity)
                return True
            # Second call — signal abort and reject
            stop_event.set()
            return False

        smu_ref = [None]
        smu = DummySMU(sleep_scale=0.0)
        smu.connect()
        smu_ref[0] = smu

        meas = LightIntensityJVMeasurement(
            sourcemeter=smu,
            intensities=_INTENSITIES,
            scan_params=_SCAN_PARAMS,
            stabilization_time_s=0,
            pre_bias_settings={"enabled": False, "voltage": 0.0, "duration": 0.0},
            light_soaking_settings={"enabled": False, "duration": 0.0},
            sample_name="AbortTest",
            output_folder=tmp_path,
            auto_single_diode_fit=False,
            on_status=lambda s: None,
            on_prompt_user=_prompt,
            on_progress=lambda i, n, s: None,
            on_measurement_done=lambda i, r: None,
            on_complete=lambda p: None,
            stop_event=stop_event,
        )
        meas.run()

        # First result completed, rest skipped
        assert meas.results[0].status == "completed"
        for r in meas.results[1:]:
            assert r.status == "skipped", f"Expected skipped, got {r.status}"

    def test_consolidated_still_written_after_abort(self, tmp_path):
        """A partial run must still produce a consolidated xlsx."""
        stop_event = threading.Event()
        stop_event.set()  # abort immediately

        meas, _ = _make_measurement(tmp_path, stop_event=stop_event)
        meas.run()

        xlsx_files = list(tmp_path.glob("*_LightIntensityJV_*.xlsx"))
        assert len(xlsx_files) == 1


# ── Auto single-diode fit ─────────────────────────────────────────────────────

class TestAutoFit:
    def test_auto_fit_populates_single_diode_fit(self, tmp_path):
        meas, _ = _make_measurement(tmp_path, auto_fit=True)
        meas.run()
        for r in meas.results:
            assert r.single_diode_fit is not None, (
                f"single_diode_fit missing for intensity={r.intensity}"
            )

    def test_sdfit_sheets_in_consolidated(self, tmp_path):
        meas, _ = _make_measurement(tmp_path, auto_fit=True)
        meas.run()
        xlsx = next(tmp_path.glob("*_LightIntensityJV_*.xlsx"))
        wb   = openpyxl.load_workbook(str(xlsx), data_only=True)
        sd_sheets = [s for s in wb.sheetnames if s.startswith("SDFit_")]
        wb.close()
        assert len(sd_sheets) == len(_INTENSITIES)


# ── Consolidated workbook content ─────────────────────────────────────────────

class TestConsolidatedContent:
    def test_config_sheet_has_correct_sample_name(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        xlsx = next(tmp_path.glob("*_LightIntensityJV_*.xlsx"))
        wb   = openpyxl.load_workbook(str(xlsx), data_only=True)
        ws   = wb["Config"]
        sample_row = next(
            (row for row in ws.iter_rows(min_row=2, max_row=25, values_only=True)
             if row and row[0] == "Sample Name"),
            None,
        )
        wb.close()
        assert sample_row is not None and sample_row[1] == "TestCell"

    def test_summary_sheet_row_count(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        xlsx = next(tmp_path.glob("*_LightIntensityJV_*.xlsx"))
        wb   = openpyxl.load_workbook(str(xlsx), data_only=True)
        ws   = wb["Summary"]
        data_rows = [
            row for row in ws.iter_rows(min_row=2, values_only=True)
            if row and row[0] is not None
        ]
        wb.close()
        assert len(data_rows) == len(_INTENSITIES)

    def test_jv_sheets_present_for_each_intensity(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        xlsx = next(tmp_path.glob("*_LightIntensityJV_*.xlsx"))
        wb   = openpyxl.load_workbook(str(xlsx), data_only=True)
        jv_sheets = [s for s in wb.sheetnames if s.startswith("JV_")]
        wb.close()
        # 2 sheets (Fwd + Rev) per completed intensity
        assert len(jv_sheets) == len(_INTENSITIES) * 2

    def test_light_ideality_sheet_present_for_4_intensities(self, tmp_path):
        meas, _ = _make_measurement(tmp_path)
        meas.run()
        xlsx = next(tmp_path.glob("*_LightIntensityJV_*.xlsx"))
        wb   = openpyxl.load_workbook(str(xlsx), data_only=True)
        assert "LightIdealityAnalysis" in wb.sheetnames
        wb.close()
