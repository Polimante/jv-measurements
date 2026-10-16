"""Lost-contact detection in illuminated J-V scans, and the SMU readback.

The failure this guards against was seen on the bench (D3C1, 2026-10-16): mid
sweep the current fell from ~17 mA/cm² to ~1e-10 A — the SMU noise floor —
and stayed there, even at V < 0 where Jsc should be. No dark diode does that,
so it is an open contact, and the software must say so instead of reporting
an FF of 1e8 %.
"""

from __future__ import annotations

import threading

import numpy as np
import openpyxl
import pytest

from app.instrument.keithley import measure_scan
from app.io.exporter import READBACK_SHEET, save_data
from app.measurement.analysis import (
    OPEN_CIRCUIT_CURRENT_A,
    detect_open_circuit,
)
from app.simulation.fake_keithley import FakeKeithley

AREA = 0.06


def _healthy_rev():
    v = np.round(np.arange(1.05, -0.1 - 1e-9, -0.01), 3)
    j = 18.2 - 18.2 * np.clip(v / 0.8, 0, None) ** 6 - 2.0 * np.clip(v - 0.8, 0, None) * 100
    return v, j


class TestDetectOpenCircuit:
    def test_healthy_scan_passes(self):
        v, j = _healthy_rev()
        assert detect_open_circuit(v, j, AREA) is None

    def test_single_point_at_voc_is_not_a_fault(self):
        # The sweep may land exactly on Voc: one ~0 point is physics.
        v, j = _healthy_rev()
        k = int(np.argmin(np.abs(j)))
        j[k] = 0.0
        assert detect_open_circuit(v, j, AREA) is None

    def test_contact_lost_mid_scan(self):
        # Shape of D3C1_morena_test_20261016_132136: drop at 0.53 V.
        v, j = _healthy_rev()
        k = int(np.where(np.isclose(v, 0.53))[0][0])
        j[k:] = 1e-6 * np.random.default_rng(0).standard_normal(len(j) - k)
        ev = detect_open_circuit(v, j, AREA)
        assert ev is not None
        assert ev.first_index == k
        assert ev.voltage == pytest.approx(0.53)
        assert ev.n_points == len(j) - k

    def test_brief_recovery_still_flagged(self):
        # The stabilized-current trace flickered back before dying for good.
        v, j = _healthy_rev()
        j[60:63] = 0.0
        j[63] = 1.6
        j[64:] = 0.0
        ev = detect_open_circuit(v, j, AREA)
        assert ev is not None and ev.first_index == 60

    def test_threshold_is_in_amperes_not_density(self):
        # 1e-9 A on 0.06 cm² is 1.67e-5 mA/cm²: just above must pass.
        v = np.linspace(0, 1, 10)
        j_just_above = np.full(10, 2.0 * OPEN_CIRCUIT_CURRENT_A * 1000 / AREA)
        assert detect_open_circuit(v, j_just_above, AREA) is None
        assert detect_open_circuit(v, j_just_above / 4, AREA) is not None


class TestMeasureScanReadback:
    def test_readback_records_every_point(self):
        k = FakeKeithley.from_mode("ideal_stable", area=AREA)
        rb = []
        v, j = measure_scan(k, -0.1, 1.0, 0.1, AREA, threading.Event(), readback=rb)
        assert len(rb) == len(v)
        for (v_set, v_meas, i_raw), vi, ji in zip(rb, v, j):
            assert v_set == pytest.approx(vi)
            assert v_meas == pytest.approx(vi, abs=1e-5)
            assert -i_raw * 1000 / AREA == pytest.approx(ji)

    def test_readback_is_optional(self):
        k = FakeKeithley.from_mode("ideal_stable", area=AREA)
        v, j = measure_scan(k, -0.1, 1.0, 0.1, AREA, threading.Event())
        assert len(v) == len(j) > 0


class TestReadbackSheet:
    def _save(self, folder, **kw):
        v = np.linspace(-0.1, 1.0, 12)
        j = 18 - 18 * np.clip(v / 0.8, 0, None) ** 6
        return save_data(
            str(folder), "Cell", v, j, v[::-1], j[::-1], 100.0,
            0.8, 18, 70, 10, 5, 1e4, 0.8, 18, 70, 10, 5, 1e4, 0.0,
            AREA, 33.3, 1.0, False, False, 0, 0, "IDN", **kw)

    def test_sheet_written_with_notes(self, tmp_path):
        rb = [(0.1 * i, 0.1 * i + 1e-4, -1e-3) for i in range(5)]
        fp = self._save(tmp_path, readback_fwd=rb, readback_rev=rb[::-1],
                        readback_notes=["Reverse scan: open contact"])
        wb = openpyxl.load_workbook(fp)
        assert wb.sheetnames == ["MeasurementData", READBACK_SHEET]
        ws = wb[READBACK_SHEET]
        assert ws.cell(2, 2).value == pytest.approx(1e-4)
        assert ws.cell(2, 3).value == pytest.approx(-1e-3)
        assert ws.cell(6, 4).value == pytest.approx(0.0)
        assert ws.cell(2, 8).value == "Reverse scan: open contact"

    def test_no_sheet_without_readback(self, tmp_path):
        fp = self._save(tmp_path)
        assert openpyxl.load_workbook(fp).sheetnames == ["MeasurementData"]

    def test_file_still_opens_in_data_viewer(self, tmp_path):
        from app.gui.data_viewer import load_jv_file
        fp = self._save(tmp_path, readback_fwd=[(0.0, 0.0, -1e-3)])
        meta = load_jv_file(fp)
        assert meta["cell_area"] == pytest.approx(AREA)
